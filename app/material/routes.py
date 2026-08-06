import os
import io
import json
from flask import render_template, request, redirect, url_for, flash, current_app, send_from_directory, session, jsonify
from flask_login import login_required
from werkzeug.utils import secure_filename
from sqlalchemy.orm import selectinload
from app.material import bp
from app import db
from app.models import Material, Category, ProjectMaterial
from app.decorators import editor_required, log_audit
from app.utils import (
    log_operation,
    get_project_materials,
    add_material_to_project,
    remove_material_from_project,
)
import openpyxl
from openpyxl.styles import Font
import qrcode


def _gen_material_code(project_id, category_id):
    """生成物资编码：三级分类编码 + 3位流水号（如 MC010601005）。

    流水号取该分类下已有材料编码末3位的最大值+1。
    """
    category = Category.query.get(category_id)
    if not category or category.level != 3 or not category.category_code:
        return None

    prefix = category.category_code
    materials = Material.query.filter(
        Material.project_id == project_id,
        Material.category_id == category_id,
        Material.code.like(f'{prefix}%')
    ).all()

    max_seq = 0
    for m in materials:
        if not m.code or len(m.code) < len(prefix):
            continue
        try:
            seq = int(m.code[len(prefix):])
            if seq > max_seq:
                max_seq = seq
        except ValueError:
            continue

    return f'{prefix}{max_seq + 1:03d}'


@bp.route('/')
@login_required
def index():
    project_id = session.get('current_project_id')
    if not project_id:
        flash('请先选择项目。', 'warning')
        return redirect(url_for('main.index'))

    page = request.args.get('page', 1, type=int)
    keyword = request.args.get('keyword', '', type=str)
    category_id = request.args.get('category_id', 0, type=int)

    # 通过 ProjectMaterial 关联表获取项目常用物资（包含公司主库 + 项目级临时物资）
    query = get_project_materials(project_id, common_only=True)
    if keyword:
        query = query.filter(Material.name.contains(keyword) | Material.code.contains(keyword))
    if category_id:
        query = query.filter(Material.category_id == category_id)

    # P2: 预加载分类信息，避免 N+1 查询
    query = query.options(selectinload(Material.category))
    pagination = query.order_by(Material.code.asc()).paginate(
        page=page, per_page=10, error_out=False
    )
    # 标记每条物资是否已加入项目常用（用于模板控制"移除常用"按钮）
    linked_ids = set(
        pm.material_id for pm in ProjectMaterial.query.filter_by(project_id=project_id).all()
    )
    for m in pagination.items:
        m._is_project_common = m.id in linked_ids

    categories = Category.query.filter_by(project_id=project_id).order_by(Category.sort_order.asc()).all()
    return render_template('material/index.html', pagination=pagination, keyword=keyword,
                           category_id=category_id, categories=categories,
                           linked_material_ids=sorted(linked_ids))


@bp.route('/create', methods=['GET', 'POST'])
@login_required
@editor_required
@log_audit(module='material', operation='新增')
def create():
    project_id = session.get('current_project_id')
    if not project_id:
        flash('请先选择项目。', 'warning')
        return redirect(url_for('main.index'))

    if request.method == 'POST':
        category_id = request.form.get('category_id', type=int) or None
        if not category_id:
            flash('请选择物资分类。', 'danger')
            return redirect(url_for('material.create'))

        category = Category.query.get(category_id)
        if not category or category.level != 3:
            flash('物资分类必须为三级分类。', 'danger')
            return redirect(url_for('material.create'))

        code = _gen_material_code(project_id, category_id)
        if not code:
            flash('物资编码生成失败，请检查分类编码是否完整。', 'danger')
            return redirect(url_for('material.create'))

        material = Material(
            project_id=project_id,
            name=request.form.get('name', '').strip(),
            code=code,
            specification=request.form.get('specification', '').strip(),
            category_id=category_id,
            unit=request.form.get('unit', '').strip(),
            remark=request.form.get('remark', '').strip(),
            # 主数据统一改造：项目级新增物资标记为 project，仅本项目可见
            source='project',
            status='active',
        )
        db.session.add(material)
        db.session.commit()
        # 项目级物资创建后自动加入项目常用表
        add_material_to_project(material.id, project_id)
        flash(f'材料创建成功，编码：{code}', 'success')
        return redirect(url_for('material.index'))

    categories = Category.query.filter_by(project_id=project_id).order_by(Category.sort_order.asc()).all()
    return render_template('material/form.html', categories=categories)


@bp.route('/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@editor_required
@log_audit(module='material', operation='编辑')
def edit(id):
    material = Material.query.get_or_404(id)
    if request.method == 'POST':
        category_id = request.form.get('category_id', type=int) or None
        if not category_id:
            flash('请选择物资分类。', 'danger')
            return redirect(url_for('material.edit', id=id))

        category = Category.query.get(category_id)
        if not category or category.level != 3:
            flash('物资分类必须为三级分类。', 'danger')
            return redirect(url_for('material.edit', id=id))

        # 编码不变，仅修改其他字段
        material.name = request.form.get('name', '').strip()
        material.specification = request.form.get('specification', '').strip()
        material.category_id = category_id
        material.unit = request.form.get('unit', '').strip()
        material.remark = request.form.get('remark', '').strip()
        db.session.commit()
        flash('材料更新成功。', 'success')
        return redirect(url_for('material.index'))

    categories = Category.query.filter_by(project_id=material.project_id).order_by(Category.sort_order.asc()).all()
    return render_template('material/form.html', material=material, categories=categories)


@bp.route('/<int:id>/delete', methods=['POST'])
@login_required
@editor_required
@log_audit(module='material', operation='删除')
def delete(id):
    material = Material.query.get_or_404(id)
    from app.models import ContractItem, StockInItem
    if ContractItem.query.filter_by(material_id=material.id).first():
        flash('该材料已关联合同明细，无法删除。', 'danger')
        return redirect(url_for('material.index'))
    if StockInItem.query.filter_by(material_id=material.id).first():
        flash('该材料已存在入库记录，无法删除。', 'danger')
        return redirect(url_for('material.index'))
    db.session.delete(material)
    db.session.commit()
    flash('材料删除成功。', 'success')
    return redirect(url_for('material.index'))


@bp.route('/<int:material_id>/remove_from_project', methods=['POST'])
@login_required
@editor_required
@log_audit(module='material', operation='移除常用')
def remove_from_project(material_id):
    """从项目常用物资移除（仅删除 ProjectMaterial 关联，不影响主库数据）"""
    project_id = session.get('current_project_id')
    if not project_id:
        flash('请先选择项目。', 'warning')
        return redirect(url_for('material.index'))

    material = Material.query.get(material_id)
    if not material:
        flash('物资不存在。', 'danger')
        return redirect(url_for('material.index'))

    # 项目级物资（source='project'）不允许仅移除常用，否则会变成游离数据
    if material.source == 'project':
        flash('项目级临时物资无法仅移除常用，请直接删除。', 'warning')
        return redirect(url_for('material.index'))

    remove_material_from_project(material_id, project_id)
    flash(f'已将物资「{material.name}」从本项目常用移除。', 'success')
    return redirect(url_for('material.index'))


@bp.route('/api/add_to_project', methods=['POST'])
@login_required
@editor_required
def api_add_to_project():
    """批量将公司库物资加入项目常用（AJAX 接口）

    请求 JSON: {"material_ids": [1, 2, 3]}
    返回 JSON: {"success": true, "added_count": N}
    """
    project_id = session.get('current_project_id')
    if not project_id:
        return jsonify({'success': False, 'message': '请先选择项目'}), 400

    payload = request.get_json(silent=True) or {}
    material_ids = payload.get('material_ids', []) or []
    if not isinstance(material_ids, list) or not material_ids:
        return jsonify({'success': False, 'message': 'material_ids 不能为空'}), 400

    added = 0
    for mid in material_ids:
        try:
            mid_int = int(mid)
        except (TypeError, ValueError):
            continue
        # 仅允许添加公司级物资
        m = Material.query.get(mid_int)
        if not m or m.source != 'company':
            continue
        add_material_to_project(mid_int, project_id)
        added += 1

    return jsonify({'success': True, 'added_count': added})


@bp.route('/api/project_materials')
@login_required
def api_project_materials():
    """API：返回当前项目常用物资列表（JSON）

    供业务单据（入库/出库/合同等）的统一物资选择弹窗使用。
    支持关键字搜索、分类过滤、分页。

    返回 JSON:
        {items: [{id, code, name, specification, unit, category_id, category_name}], total, page, pages}
    """
    from sqlalchemy import or_

    project_id = session.get('current_project_id')
    if not project_id:
        return jsonify({'items': [], 'total': 0, 'page': 1, 'pages': 0})

    q = request.args.get('q', '', type=str)
    category_id = request.args.get('category_id', 0, type=int)
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 20, type=int)
    # 限制最大返回数量，避免一次性拉取过多
    per_page = min(max(per_page, 1), 100)

    query = get_project_materials(project_id, common_only=True)
    if q:
        query = query.filter(
            or_(Material.name.contains(q), Material.code.contains(q))
        )
    if category_id:
        query = query.filter(Material.category_id == category_id)

    # P2: 预加载分类信息，避免 N+1 查询
    query = query.options(selectinload(Material.category))
    pagination = query.order_by(Material.code.asc()).paginate(
        page=page, per_page=per_page, error_out=False
    )

    items = []
    for m in pagination.items:
        items.append({
            'id': m.id,
            'code': m.code or '',
            'name': m.name,
            'specification': m.specification or '',
            'unit': m.unit or '',
            'category_id': m.category_id,
            'category_name': m.category.name if m.category else '',
        })

    return jsonify({
        'items': items,
        'total': pagination.total,
        'page': pagination.page,
        'pages': pagination.pages,
    })


@bp.route('/download_template')
@login_required
def download_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '材料导入模板'
    headers = ['物资名称', '规格型号', '物资分类', '单位', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(['示例钢筋', 'HRB400 Φ12', 'MC010601', '吨', '用于主体'])

    template_dir = os.path.join(current_app.root_path, 'static', 'uploads')
    os.makedirs(template_dir, exist_ok=True)
    template_path = os.path.join(template_dir, '材料导入模板.xlsx')
    wb.save(template_path)
    return send_from_directory(template_dir, '材料导入模板.xlsx', as_attachment=True)


@bp.route('/import', methods=['POST'])
@login_required
@editor_required
@log_audit(module='material', operation='批量导入')
def import_excel():
    """批量导入材料，支持错误校验和提示。

    导入要求：
    - 表头：物资名称、规格型号、物资分类、单位、备注
    - 物资分类必须填写三级分类的编码（category_code），系统据此自动生成物资编码
    - 编码由系统自动生成（三级分类编码 + 3位流水号）
    """
    project_id = session.get('current_project_id')
    if not project_id:
        flash('请先选择项目。', 'warning')
        return redirect(url_for('main.index'))

    file = request.files.get('file')
    if not file:
        flash('请选择文件。', 'danger')
        return redirect(url_for('material.index'))

    try:
        wb = openpyxl.load_workbook(file)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        success_count = 0
        errors = []

        for idx, row in enumerate(rows, start=2):  # 从第2行开始（第1行是表头）
            if not row or not row[0]:
                continue

            # 数据校验
            name = str(row[0]).strip() if row[0] else ''
            if not name:
                errors.append(f'第{idx}行：物资名称不能为空')
                continue

            spec = str(row[1]).strip() if len(row) > 1 and row[1] else ''
            cat_code = str(row[2]).strip() if len(row) > 2 and row[2] else ''
            unit = str(row[3]).strip() if len(row) > 3 and row[3] else ''
            remark = str(row[4]).strip() if len(row) > 4 and row[4] else ''

            # 检查单位是否填写
            if not unit:
                errors.append(f'第{idx}行：单位不能为空')
                continue

            # 检查是否已存在同名材料
            existing = Material.query.filter_by(project_id=project_id, name=name).first()
            if existing:
                errors.append(f'第{idx}行：物资名称"{name}"已存在')
                continue

            # 处理分类：必须为三级分类编码
            if not cat_code:
                errors.append(f'第{idx}行：物资分类编码不能为空（请填写三级分类编码）')
                continue

            category = Category.query.filter_by(
                project_id=project_id, category_code=cat_code
            ).first()
            if not category:
                errors.append(f'第{idx}行：物资分类编码"{cat_code}"不存在')
                continue
            if category.level != 3:
                errors.append(f'第{idx}行：物资分类"{cat_code}"不是三级分类')
                continue

            # 自动生成物资编码
            code = _gen_material_code(project_id, category.id)
            if not code:
                errors.append(f'第{idx}行：物资编码生成失败')
                continue

            material = Material(
                project_id=project_id,
                name=name,
                code=code,
                specification=spec,
                category_id=category.id,
                unit=unit,
                remark=remark,
                # 主数据统一改造：项目级导入物资标记为 project，自动加入项目常用
                source='project',
                status='active',
            )
            db.session.add(material)
            db.session.flush()  # 取 material.id
            # 加入项目常用表（与新增逻辑保持一致）
            existing_link = ProjectMaterial.query.filter_by(
                project_id=project_id, material_id=material.id
            ).first()
            if not existing_link:
                db.session.add(ProjectMaterial(
                    project_id=project_id,
                    material_id=material.id,
                    is_common=True,
                    sort=0,
                ))
            success_count += 1

        if success_count > 0:
            db.session.commit()

        if errors:
            flash(f'成功导入 {success_count} 条，{len(errors)} 条失败：\n' + '\n'.join(errors[:5]), 'warning')
        else:
            flash(f'成功导入 {success_count} 条材料数据。', 'success')

        log_operation('导入', module='材料管理', description=f'批量导入材料 {success_count} 条')

    except Exception as e:
        db.session.rollback()
        flash(f'导入失败：{str(e)}', 'danger')

    return redirect(url_for('material.index'))


@bp.route('/api/qr_generate/<int:material_id>')
@login_required
def qr_generate(material_id):
    """为指定材料生成二维码 PNG 图片。"""
    project_id = session.get('current_project_id')
    if not project_id:
        return jsonify({'error': '请先选择项目。'}), 400

    material = Material.query.filter_by(id=material_id, project_id=project_id).first_or_404()
    data = json.dumps({
        'project_id': project_id,
        'material_id': material.id,
        'code': material.code,
        'name': material.name
    }, ensure_ascii=False)

    qr_dir = os.path.join(current_app.root_path, 'static', 'qr_codes')
    os.makedirs(qr_dir, exist_ok=True)
    filename = f'{material.code}_{material.id}.png'
    filepath = os.path.join(qr_dir, filename)

    img = qrcode.make(data)
    img.save(filepath)

    rel_path = f'/static/qr_codes/{filename}'
    return jsonify({'path': rel_path, 'filename': filename})


@bp.route('/api/qr_batch', methods=['POST'])
@login_required
def qr_batch():
    """批量生成二维码。"""
    project_id = session.get('current_project_id')
    if not project_id:
        return jsonify({'error': '请先选择项目。'}), 400

    payload = request.get_json() or {}
    material_ids = payload.get('material_ids', [])
    if not material_ids:
        return jsonify({'error': 'material_ids 不能为空。'}), 400

    qr_dir = os.path.join(current_app.root_path, 'static', 'qr_codes')
    os.makedirs(qr_dir, exist_ok=True)

    results = []
    for mid in material_ids:
        material = Material.query.filter_by(id=mid, project_id=project_id).first()
        if not material:
            results.append({'material_id': mid, 'error': '材料不存在'})
            continue

        data = json.dumps({
            'project_id': project_id,
            'material_id': material.id,
            'code': material.code,
            'name': material.name
        }, ensure_ascii=False)

        filename = f'{material.code}_{material.id}.png'
        filepath = os.path.join(qr_dir, filename)
        img = qrcode.make(data)
        img.save(filepath)

        results.append({
            'material_id': material.id,
            'path': f'/static/qr_codes/{filename}',
            'filename': filename
        })

    return jsonify({'results': results})


@bp.route('/api/scan_lookup')
@login_required
def scan_lookup():
    """根据物资编码查询当前项目常用材料信息。

    主数据统一改造后，扫码查询限定在项目常用物资范围内（与表单选择器一致）。
    """
    project_id = session.get('current_project_id')
    if not project_id:
        return jsonify({'error': '请先选择项目。'}), 400

    code = request.args.get('code', '').strip()
    if not code:
        return jsonify({'error': 'code 参数不能为空。'}), 400

    material = get_project_materials(project_id, common_only=True).filter(
        Material.code == code
    ).first()
    if not material:
        return jsonify({'error': '未找到该物资编码对应的材料（请先将其加入本项目常用物资）。'}), 404

    return jsonify({
        'id': material.id,
        'code': material.code,
        'name': material.name,
        'specification': material.specification or '',
        'unit': material.unit or '',
        'category_id': material.category_id,
        'category_name': material.category.name if material.category else ''
    })
