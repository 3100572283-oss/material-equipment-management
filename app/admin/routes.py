import os
import shutil
from datetime import datetime
from flask import (render_template, request, redirect, url_for, flash,
                   send_file, session, current_app, jsonify)
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename
from app.admin import bp
from app.decorators import admin_required, log_audit
from app.models import User, SystemConfig, OperationLog, SysDept, SysRole, Project, SysUserProject
from app import db


# ============== 用户管理 ==============

@bp.route('/users')
@login_required
@admin_required
def users():
    users = User.query.order_by(User.created_at.desc()).all()
    depts = SysDept.query.filter_by(status=True).order_by(SysDept.sort.asc(), SysDept.created_at.asc()).all()
    roles = SysRole.query.filter_by(status=True).order_by(SysRole.sort).all()
    from app.utils import get_config
    default_password = get_config('default_password', 'Abc@123456')
    return render_template('admin/users.html', users=users, depts=depts, roles=roles, default_password=default_password)


@bp.route('/users/<int:id>/projects', methods=['GET', 'POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='配置用户项目权限')
def user_projects(id):
    """用户项目权限配置"""
    user = User.query.get_or_404(id)
    all_projects = Project.query.filter_by(is_archived=False).order_by(Project.created_at.desc()).all()
    if request.method == 'POST':
        project_ids = request.form.getlist('project_ids', type=int)
        main_project_id = request.form.get('main_project_id', type=int)

        # 全部数据权限用户不可配置（自动拥有所有项目）
        if user.get_data_scope() == 'all' or user.is_admin():
            flash('该用户拥有全部数据权限，自动可访问所有项目，无需单独配置', 'info')
            return redirect(url_for('admin.user_projects', id=id))

        # 删除旧关联
        SysUserProject.query.filter_by(user_id=user.id).delete()
        db.session.flush()

        # 创建新关联
        for pid in project_ids:
            is_main = (pid == main_project_id)
            up = SysUserProject(user_id=user.id, project_id=pid, is_main=is_main)
            db.session.add(up)

        db.session.commit()
        flash('用户项目权限已更新', 'success')
        return redirect(url_for('admin.user_projects', id=id))

    user_project_ids = [up.project_id for up in user.user_projects]
    user_main_project_id = None
    for up in user.user_projects:
        if up.is_main:
            user_main_project_id = up.project_id
            break
    return render_template('admin/user_projects.html', user=user, all_projects=all_projects,
                           user_project_ids=user_project_ids, user_main_project_id=user_main_project_id)


@bp.route('/users/create', methods=['GET', 'POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='新增用户')
def create_user():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        role_id = request.form.get('role_id', type=int)
        dept_id = request.form.get('dept_id', type=int)
        project_id = request.form.get('project_id', type=int)
        name = request.form.get('name', '').strip()

        if not username or not password:
            flash('用户名和密码不能为空', 'error')
            return redirect(url_for('admin.create_user'))

        if User.query.filter_by(username=username).first():
            flash('用户名已存在', 'error')
            return redirect(url_for('admin.create_user'))

        from werkzeug.security import generate_password_hash
        
        auto_project_ids = []
        auto_main_project_id = None
        
        if dept_id:
            dept = SysDept.query.get(dept_id)
            if dept and dept.dept_type == 'project' and dept.project_id:
                auto_project_ids.append(dept.project_id)
                auto_main_project_id = dept.project_id
        
        form_project_ids = request.form.getlist('project_ids', type=int)
        form_main_project_id = request.form.get('main_project_id', type=int)
        
        all_project_ids = list(set(auto_project_ids + form_project_ids))
        
        if form_main_project_id:
            final_main_project_id = form_main_project_id
        elif auto_main_project_id:
            final_main_project_id = auto_main_project_id
        elif all_project_ids:
            final_main_project_id = all_project_ids[0]
        else:
            final_main_project_id = None
        
        user = User(
            username=username,
            password_hash=generate_password_hash(password, method='pbkdf2:sha256'),
            role='viewer',
            role_id=role_id,
            dept_id=dept_id,
            project_id=final_main_project_id,
            name=name or None,
            department=request.form.get('department', '').strip() or None,
            email=request.form.get('email', '').strip() or None,
            phone=request.form.get('phone', '').strip() or None,
        )
        db.session.add(user)
        db.session.flush()
        
        for pid in all_project_ids:
            is_main = (pid == final_main_project_id)
            up = SysUserProject(user_id=user.id, project_id=pid, is_main=is_main)
            db.session.add(up)
        
        db.session.commit()
        flash('用户创建成功', 'success')
        return redirect(url_for('admin.users'))

    depts = SysDept.query.filter_by(status=True).order_by(SysDept.sort.asc(), SysDept.created_at.asc()).all()
    roles = SysRole.query.filter_by(status=True).order_by(SysRole.sort).all()
    projects = Project.query.filter_by(is_archived=False).order_by(Project.code).all()

    dept_map = {d.id: d for d in depts}
    def get_dept_path(d):
        path = []
        current = d
        while current:
            path.insert(0, current.dept_name)
            if current.parent_id and current.parent_id in dept_map:
                current = dept_map[current.parent_id]
            else:
                current = None
        return ' / '.join(path)
    for d in depts:
        d._path = get_dept_path(d)

    dept_project_map = {}
    for d in depts:
        if d.dept_type == 'project' and d.project_id:
            dept_project_map[d.id] = d.project_id

    # 支持从查询参数预填部门
    preset_dept_id = request.args.get('dept_id', type=int)

    return render_template('admin/user_form.html', user=None, depts=depts, roles=roles,
                           projects=projects, dept_project_map=dept_project_map,
                           preset_dept_id=preset_dept_id)


@bp.route('/users/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='编辑用户')
def edit_user(id):
    user = User.query.get_or_404(id)
    if request.method == 'POST':
        role_id = request.form.get('role_id', type=int)
        dept_id = request.form.get('dept_id', type=int)
        name = request.form.get('name', '').strip()
        user.role_id = role_id
        user.dept_id = dept_id
        user.name = name or None
        user.department = request.form.get('department', '').strip() or None
        user.email = request.form.get('email', '').strip() or None
        user.phone = request.form.get('phone', '').strip() or None

        # 重置密码
        new_password = request.form.get('new_password', '')
        if new_password:
            from werkzeug.security import generate_password_hash
            user.password_hash = generate_password_hash(new_password, method='pbkdf2:sha256')

        # ===== 更新项目关联 =====
        auto_project_ids = []
        auto_main_project_id = None

        if dept_id:
            dept = SysDept.query.get(dept_id)
            if dept and dept.dept_type == 'project' and dept.project_id:
                auto_project_ids.append(dept.project_id)
                auto_main_project_id = dept.project_id

        form_project_ids = request.form.getlist('project_ids', type=int)
        form_main_project_id = request.form.get('main_project_id', type=int)

        all_project_ids = list(set(auto_project_ids + form_project_ids))

        if form_main_project_id:
            final_main_project_id = form_main_project_id
        elif auto_main_project_id:
            final_main_project_id = auto_main_project_id
        elif all_project_ids:
            final_main_project_id = all_project_ids[0]
        else:
            final_main_project_id = None

        # 更新主项目
        user.project_id = final_main_project_id

        # 删除旧关联，重建新关联
        SysUserProject.query.filter_by(user_id=user.id).delete()
        for pid in all_project_ids:
            is_main = (pid == final_main_project_id)
            up = SysUserProject(user_id=user.id, project_id=pid, is_main=is_main)
            db.session.add(up)

        db.session.commit()
        flash('用户信息更新成功', 'success')
        return redirect(url_for('admin.users'))

    depts = SysDept.query.filter_by(status=True).order_by(SysDept.sort.asc(), SysDept.created_at.asc()).all()
    roles = SysRole.query.filter_by(status=True).order_by(SysRole.sort).all()
    projects = Project.query.filter_by(is_archived=False).order_by(Project.code).all()

    dept_map = {d.id: d for d in depts}
    def get_dept_path(d):
        path = []
        current = d
        while current:
            path.insert(0, current.dept_name)
            if current.parent_id and current.parent_id in dept_map:
                current = dept_map[current.parent_id]
            else:
                current = None
        return ' / '.join(path)
    for d in depts:
        d._path = get_dept_path(d)

    dept_project_map = {}
    for d in depts:
        if d.dept_type == 'project' and d.project_id:
            dept_project_map[d.id] = d.project_id

    from app.utils import get_config
    default_password = get_config('default_password', 'Abc@123456')
    return render_template('admin/user_form.html', user=user, depts=depts, roles=roles,
                           projects=projects, dept_project_map=dept_project_map,
                           default_password=default_password, preset_dept_id=None)


@bp.route('/users/<int:id>/reset_password', methods=['GET', 'POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='重置密码')
def reset_password(id):
    user = User.query.get_or_404(id)
    from app.utils import get_config
    default_pwd = get_config('default_password', 'Abc@123456')

    if request.method == 'POST':
        new_password = request.form.get('new_password', '')
        confirm_password = request.form.get('confirm_password', '')
        force_change = request.form.get('force_change') == '1'

        if not new_password:
            flash('新密码不能为空', 'error')
            return redirect(url_for('admin.reset_password', id=id))
        if new_password != confirm_password:
            flash('两次输入的密码不一致', 'error')
            return redirect(url_for('admin.reset_password', id=id))
        if len(new_password) < 6:
            flash('密码长度不能少于6位', 'error')
            return redirect(url_for('admin.reset_password', id=id))

        from werkzeug.security import generate_password_hash
        user.password_hash = generate_password_hash(new_password, method='pbkdf2:sha256')
        user.must_change_password = force_change
        db.session.commit()

        # 操作日志：记录重置密码详情
        from app.utils import log_operation
        log_operation(
            '重置密码',
            module='用户管理',
            description=f'管理员 {current_user.username} 重置用户 {user.username} 的密码'
                         + ('，强制下次登录修改密码' if force_change else '')
        )

        flash(f'用户 {user.username} 密码已重置成功', 'success')
        return redirect(url_for('admin.users'))

    return render_template('admin/reset_password.html', user=user, default_password=default_pwd)


@bp.route('/users/<int:id>/delete', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='删除用户')
def delete_user(id):
    user = User.query.get_or_404(id)
    if user.id == current_user.id:
        flash('不能删除当前登录用户', 'error')
        return redirect(url_for('admin.users'))
    db.session.delete(user)
    db.session.commit()
    flash('用户已删除', 'success')
    return redirect(url_for('admin.users'))


@bp.route('/users/import_template')
@login_required
@admin_required
def import_user_template():
    """下载用户批量导入模板"""
    from io import BytesIO
    from openpyxl import Workbook
    from flask import make_response
    wb = Workbook()
    ws = wb.active
    ws.title = '用户导入模板'
    headers = ['用户名*', '姓名*', '部门编码', '角色编码', '手机号', '邮箱', '初始密码']
    ws.append(headers)
    # 示例行
    ws.append(['zhangsan', '张三', 'DEPT01', 'viewer', '13800138000', 'zhangsan@example.com', '123456'])
    ws.append(['lisi', '李四', 'DEPT02', 'editor', '13800138001', 'lisi@example.com', ''])

    # 说明sheet
    ws2 = wb.create_sheet('填写说明')
    ws2['A1'] = '填写说明'
    ws2['A1'].font = ws2['A1'].font.copy(bold=True)
    notes = [
        '1. 带*号的为必填项',
        '2. 部门编码：请在部门管理中查看对应编码',
        '3. 角色编码：请在角色管理中查看对应编码（如 super_admin, admin, editor, viewer）',
        '4. 初始密码：留空则使用系统默认密码（123456）',
        '5. 用户名不能重复，长度4-20位',
        '6. 导入成功后，新用户首次登录建议修改密码',
    ]
    for i, note in enumerate(notes, start=3):
        ws2.cell(row=i, column=1, value=note)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    resp = make_response(output.getvalue())
    resp.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    resp.headers['Content-Disposition'] = 'attachment; filename=user_import_template.xlsx'
    return resp


@bp.route('/users/import', methods=['GET', 'POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='批量导入用户')
def import_users():
    """用户批量导入"""
    if request.method == 'POST':
        from openpyxl import load_workbook
        from werkzeug.security import generate_password_hash
        from datetime import datetime
        import os

        file = request.files.get('file')
        if not file or not file.filename:
            flash('请选择要导入的Excel文件', 'error')
            return redirect(url_for('admin.import_users'))

        # 检查文件类型
        ext = os.path.splitext(file.filename)[1].lower()
        if ext not in ['.xlsx', '.xls']:
            flash('仅支持 .xlsx 或 .xls 格式的Excel文件', 'error')
            return redirect(url_for('admin.import_users'))

        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            flash(f'文件读取失败：{str(e)}', 'error')
            return redirect(url_for('admin.import_users'))

        # 读取数据（跳过表头）
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        if not rows:
            flash('文件中没有数据', 'error')
            return redirect(url_for('admin.import_users'))

        # 预加载部门和角色映射
        dept_map = {d.dept_code: d.id for d in SysDept.query.all() if d.dept_code}
        role_map = {r.role_code: r.id for r in SysRole.query.all()}
        default_password = '123456'

        success_count = 0
        error_rows = []
        existing_usernames = {u.username for u in User.query.all()}

        for idx, row in enumerate(rows, start=2):
            # 空行跳过
            if not row or all(c is None or str(c).strip() == '' for c in row):
                continue

            username = str(row[0]).strip() if row[0] else ''
            name = str(row[1]).strip() if row[1] else ''
            dept_code = str(row[2]).strip() if row[2] else ''
            role_code = str(row[3]).strip() if row[3] else ''
            phone = str(row[4]).strip() if row[4] else ''
            email = str(row[5]).strip() if row[5] else ''
            password = str(row[6]).strip() if row[6] else ''

            errors = []

            # 校验必填
            if not username:
                errors.append('用户名为空')
            elif len(username) < 4 or len(username) > 20:
                errors.append('用户名长度应为4-20位')
            elif username in existing_usernames:
                errors.append('用户名已存在')

            if not name:
                errors.append('姓名为空')

            # 校验部门
            dept_id = None
            if dept_code:
                if dept_code in dept_map:
                    dept_id = dept_map[dept_code]
                else:
                    errors.append(f'部门编码「{dept_code}」不存在')

            # 校验角色
            role_id = None
            if role_code:
                if role_code in role_map:
                    role_id = role_map[role_code]
                else:
                    errors.append(f'角色编码「{role_code}」不存在')

            if errors:
                error_rows.append({
                    'row': idx,
                    'username': username,
                    'name': name,
                    'errors': errors
                })
                continue

            # 创建用户
            try:
                pwd = password or default_password
                user = User(
                    username=username,
                    password_hash=generate_password_hash(pwd, method='pbkdf2:sha256'),
                    role='viewer',
                    role_id=role_id,
                    dept_id=dept_id,
                    name=name,
                    phone=phone or None,
                    email=email or None,
                )
                db.session.add(user)
                existing_usernames.add(username)
                success_count += 1
            except Exception as e:
                error_rows.append({
                    'row': idx,
                    'username': username,
                    'name': name,
                    'errors': [f'创建失败：{str(e)}']
                })

        db.session.commit()

        return render_template('admin/user_import_result.html',
                               success_count=success_count,
                               error_rows=error_rows,
                               total=len(rows))

    return render_template('admin/user_import.html')


# ============== 系统配置 ==============

@bp.route('/config')
@login_required
@admin_required
def config_list():
    configs = SystemConfig.query.order_by(SystemConfig.config_key).all()
    return render_template('admin/config.html', configs=configs)


@bp.route('/config/<int:id>/edit', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='保存配置')
def edit_config(id):
    config = SystemConfig.query.get_or_404(id)
    config.config_value = request.form.get('config_value', '')
    db.session.commit()
    from app.utils import ConfigCache
    ConfigCache.clear()
    flash('配置已更新', 'success')
    return redirect(url_for('admin.config_list'))


# ============== 数据备份 ==============

BACKUP_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'backups')


def ensure_backup_dir():
    if not os.path.exists(BACKUP_DIR):
        os.makedirs(BACKUP_DIR)


@bp.route('/backup')
@login_required
@admin_required
def backup_index():
    ensure_backup_dir()
    backups = []
    for fname in sorted(os.listdir(BACKUP_DIR), reverse=True):
        fpath = os.path.join(BACKUP_DIR, fname)
        if os.path.isfile(fpath) and fname.endswith('.db'):
            size = os.path.getsize(fpath)
            backups.append({
                'filename': fname,
                'size': f'{size / 1024:.1f} KB' if size < 1024 * 1024 else f'{size / (1024 * 1024):.2f} MB',
                'created_at': datetime.fromtimestamp(os.path.getctime(fpath)).strftime('%Y-%m-%d %H:%M:%S')
            })
    return render_template('admin/backup.html', backups=backups)


@bp.route('/backup/now')
@login_required
@admin_required
@log_audit(module='admin', operation='备份')
def backup_now():
    ensure_backup_dir()
    db_path = current_app.config['SQLALCHEMY_DATABASE_URI'].replace('sqlite:///', '')
    if not os.path.isabs(db_path):
        db_path = os.path.join(current_app.root_path, '..', db_path)
        db_path = os.path.abspath(db_path)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    backup_name = f'backup_{timestamp}.db'
    backup_path = os.path.join(BACKUP_DIR, backup_name)

    shutil.copy2(db_path, backup_path)
    flash(f'备份成功：{backup_name}', 'success')
    return redirect(url_for('admin.backup_index'))


@bp.route('/backup/download/<filename>')
@login_required
@admin_required
def download_backup(filename):
    fpath = os.path.join(BACKUP_DIR, secure_filename(filename))
    if os.path.exists(fpath):
        return send_file(fpath, as_attachment=True)
    flash('文件不存在', 'error')
    return redirect(url_for('admin.backup_index'))


@bp.route('/backup/delete/<filename>', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='删除备份')
def delete_backup(filename):
    fpath = os.path.join(BACKUP_DIR, secure_filename(filename))
    if os.path.exists(fpath):
        os.remove(fpath)
        flash('备份已删除', 'success')
    else:
        flash('文件不存在', 'error')
    return redirect(url_for('admin.backup_index'))


@bp.route('/backup/restore', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='还原')
def restore_backup():
    file = request.files.get('backup_file')
    if not file or not file.filename.endswith('.db'):
        flash('请上传有效的 .db 备份文件', 'error')
        return redirect(url_for('admin.backup_index'))

    # 先自动备份当前数据
    ensure_backup_dir()
    db_path = current_app.config['SQLALCHEMY_DATABASE_URI'].replace('sqlite:///', '')
    if not os.path.isabs(db_path):
        db_path = os.path.join(current_app.root_path, '..', db_path)
        db_path = os.path.abspath(db_path)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    auto_backup = os.path.join(BACKUP_DIR, f'auto_backup_before_restore_{timestamp}.db')
    shutil.copy2(db_path, auto_backup)

    # 恢复数据
    temp_path = os.path.join(BACKUP_DIR, 'temp_restore.db')
    file.save(temp_path)
    shutil.copy2(temp_path, db_path)
    os.remove(temp_path)

    flash('数据恢复成功，请重新登录', 'success')
    return redirect(url_for('auth.logout'))


# ============== 操作日志 ==============

@bp.route('/logs')
@login_required
@admin_required
def logs():
    page = request.args.get('page', 1, type=int)
    query = OperationLog.query.order_by(OperationLog.created_at.desc())
    pagination = query.paginate(page=page, per_page=20, error_out=False)
    return render_template('admin/logs.html', pagination=pagination)


# ============== 回收站 ==============

@bp.route('/recycle_bin')
@login_required
@admin_required
def recycle_bin():
    from sqlalchemy import text
    from datetime import datetime, timedelta
    from app.utils import get_config
    retention = int(get_config('recycle_bin_retention_days', '30'))
    cutoff = datetime.now() - timedelta(days=retention)
    tables = [
        ('contracts', '合同', 'contract_no', 'name'),
        ('stock_ins', '入库单', 'stock_in_no', 'supplier_id'),
        ('stock_outs', '出库单', 'stock_out_no', 'work_number'),
        ('suppliers', '供应商', 'name', 'contact_person'),
        ('materials', '物资', 'name', 'specification'),
        ('payments', '付款单', 'payment_no', 'contract_id'),
        ('equipment', '设备', 'name', 'code'),
        ('turnover_material', '周转材', 'name', 'code'),
    ]
    items = []
    for tbl, label, col1, col2 in tables:
        try:
            rows = db.session.execute(text(f"SELECT id, {col1}, {col2}, deleted_at FROM {tbl} WHERE is_deleted = 1 AND deleted_at >= :cutoff ORDER BY deleted_at DESC"), {'cutoff': cutoff}).fetchall()
            for r in rows:
                items.append({'table': tbl, 'label': label, 'id': r[0], 'name': str(r[1] or ''), 'extra': str(r[2] or ''), 'deleted_at': r[3]})
        except Exception:
            pass
    return render_template('admin/recycle_bin.html', items=items)

@bp.route('/recycle_bin/<table>/<int:id>/restore', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='恢复回收站数据')
def restore_item(table, id):
    from sqlalchemy import text
    allowed = ['contracts', 'stock_ins', 'stock_outs', 'suppliers', 'materials', 'payments', 'equipment', 'turnover_material']
    if table not in allowed:
        flash('不支持的表', 'danger')
        return redirect(url_for('admin.recycle_bin'))
    db.session.execute(text(f"UPDATE {table} SET is_deleted = 0, deleted_at = NULL WHERE id = :id"), {'id': id})
    db.session.commit()
    flash('数据已恢复', 'success')
    return redirect(url_for('admin.recycle_bin'))

@bp.route('/recycle_bin/<table>/<int:id>/purge', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='彻底删除')
def purge_item(table, id):
    from sqlalchemy import text
    allowed = ['contracts', 'stock_ins', 'stock_outs', 'suppliers', 'materials', 'payments', 'equipment', 'turnover_material']
    if table not in allowed:
        flash('不支持的表', 'danger')
        return redirect(url_for('admin.recycle_bin'))
    db.session.execute(text(f"DELETE FROM {table} WHERE id = :id AND is_deleted = 1"), {'id': id})
    db.session.commit()
    flash('数据已永久删除', 'success')
    return redirect(url_for('admin.recycle_bin'))


# ============== 数据归档 ==============

@bp.route('/archive')
@login_required
@admin_required
def archive_index():
    from app.models import Project
    archived = Project.query.filter_by(is_archived=True).order_by(Project.created_at.desc()).all()
    active = Project.query.filter_by(is_archived=False).order_by(Project.created_at.desc()).all()
    return render_template('admin/archive.html', archived=archived, active=active)

@bp.route('/archive/<int:project_id>/toggle', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='归档/取消归档')
def toggle_archive(project_id):
    from app.models import Project
    from flask import session
    project = Project.query.get_or_404(project_id)
    project.is_archived = not project.is_archived
    db.session.commit()
    if session.get('current_project_id') == project_id and project.is_archived:
        session.pop('current_project_id', None)
        session.pop('current_project_name', None)
    flash(f'项目已{"归档" if project.is_archived else "取消归档"}', 'success')
    return redirect(url_for('admin.archive_index'))


# ============== 审计日志（增强版） ==============

@bp.route('/audit_logs')
@login_required
@admin_required
def audit_logs():
    from app.models import SysOperationLog
    page = request.args.get('page', 1, type=int)
    username = request.args.get('username', '', type=str)
    module = request.args.get('module', '', type=str)
    operation = request.args.get('operation', '', type=str)
    status = request.args.get('status', '', type=str)
    start_date = request.args.get('start_date', '', type=str)
    end_date = request.args.get('end_date', '', type=str)
    query = SysOperationLog.query
    if username: query = query.filter(SysOperationLog.username.contains(username))
    if module: query = query.filter(SysOperationLog.module == module)
    if operation: query = query.filter(SysOperationLog.operation == operation)
    if status: query = query.filter(SysOperationLog.status == status)
    if start_date:
        from datetime import datetime
        query = query.filter(SysOperationLog.operation_time >= datetime.strptime(start_date, '%Y-%m-%d'))
    if end_date:
        from datetime import datetime, timedelta
        query = query.filter(SysOperationLog.operation_time < datetime.strptime(end_date, '%Y-%m-%d') + timedelta(days=1))
    pagination = query.order_by(SysOperationLog.operation_time.desc()).paginate(page=page, per_page=20, error_out=False)
    return render_template('admin/audit_logs.html', pagination=pagination)

@bp.route('/audit_logs/<int:id>')
@login_required
@admin_required
def audit_log_detail(id):
    from app.models import SysOperationLog, DataChangeLog
    log = SysOperationLog.query.get_or_404(id)
    import json
    params = json.loads(log.params) if log.params else {}
    changes = json.loads(log.changes) if log.changes else {}
    change_logs = DataChangeLog.query.filter_by(record_id=log.biz_id).all() if log.biz_id else []
    return render_template('admin/audit_log_detail.html', log=log, params=params, changes=changes)

@bp.route('/audit_logs/export')
@login_required
@admin_required
def export_audit_logs():
    from app.models import SysOperationLog
    from app.utils import export_to_excel
    from flask import Response
    logs = SysOperationLog.query.order_by(SysOperationLog.operation_time.desc()).limit(5000).all()
    headers = ['操作时间', '操作人', '模块', '操作', '业务ID', '请求方法', 'IP', '状态', '耗时(ms)']
    rows = [[l.operation_time.strftime('%Y-%m-%d %H:%M:%S') if l.operation_time else '', l.username or '', l.module or '', l.operation or '', l.biz_id or '', l.method or '', l.ip_address or '', l.status or '', l.cost_time or 0] for l in logs]
    data = export_to_excel(headers, rows, '审计日志', 'audit_logs.xlsx')
    return Response(data, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': 'attachment; filename=audit_logs.xlsx'})


# ============== 通知配置 ==============

@bp.route('/notify_config', methods=['GET', 'POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='保存通知配置')
def notify_config():
    from app.models import SystemConfig, NotificationLog
    if request.method == 'POST':
        configs = ['notify_dingtalk_enabled', 'notify_dingtalk_webhook', 'notify_dingtalk_secret',
                   'notify_wechat_enabled', 'notify_wechat_webhook',
                   'notify_feishu_enabled', 'notify_feishu_webhook',
                   'notify_email_enabled', 'notify_email_smtp_host', 'notify_email_smtp_port',
                   'notify_email_account', 'notify_email_password', 'notify_email_recipients',
                   'notify_title_prefix']
        for key in configs:
            val = request.form.get(key, '')
            c = SystemConfig.query.filter_by(config_key=key).first()
            if c:
                c.config_value = val
            else:
                db.session.add(SystemConfig(config_key=key, config_value=val))
        db.session.commit()
        from app.utils import ConfigCache
        ConfigCache.clear()
        flash('通知配置保存成功', 'success')
        return redirect(url_for('admin.notify_config'))
    configs = {c.config_key: c.config_value for c in SystemConfig.query.all()}
    recent_logs = NotificationLog.query.order_by(NotificationLog.created_at.desc()).limit(10).all()
    return render_template('admin/notify_config.html', configs=configs, recent_logs=recent_logs)

@bp.route('/notify_test', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='测试通知')
def notify_test():
    channel = request.form.get('channel', 'dingtalk')
    from app.notification_service import push_notification
    results = push_notification('测试消息', f'这是一条来自物资系统的测试消息，发送时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', [channel])
    if results and results[0][1]:
        flash(f'{channel} 测试发送成功', 'success')
    else:
        err = results[0][2] if results else '未知错误'
        flash(f'{channel} 测试发送失败: {err}', 'danger')
    return redirect(url_for('admin.notify_config'))


# ============== 登录日志 ==============

@bp.route('/login_logs')
@login_required
@admin_required
def login_logs():
    from app.models import LoginLog
    page = request.args.get('page', 1, type=int)
    username = request.args.get('username', '', type=str)
    status = request.args.get('status', '', type=str)
    start_date = request.args.get('start_date', '', type=str)
    end_date = request.args.get('end_date', '', type=str)

    query = LoginLog.query
    if username:
        query = query.filter(LoginLog.username.contains(username))
    if status:
        query = query.filter_by(status=status)
    if start_date:
        query = query.filter(LoginLog.login_time >= datetime.strptime(start_date, '%Y-%m-%d'))
    if end_date:
        query = query.filter(LoginLog.login_time < datetime.strptime(end_date, '%Y-%m-%d').replace(hour=23, minute=59, second=59))

    pagination = query.order_by(LoginLog.login_time.desc()).paginate(page=page, per_page=20, error_out=False)
    return render_template('admin/login_logs.html', pagination=pagination, username=username, status=status, start_date=start_date, end_date=end_date)


@bp.route('/login_logs/export')
@login_required
@admin_required
def export_login_logs():
    from app.models import LoginLog
    from app.utils import export_to_excel
    username = request.args.get('username', '', type=str)
    status = request.args.get('status', '', type=str)

    query = LoginLog.query
    if username:
        query = query.filter(LoginLog.username.contains(username))
    if status:
        query = query.filter_by(status=status)

    logs = query.order_by(LoginLog.login_time.desc()).all()
    headers = ['登录时间', '用户名', 'IP地址', '浏览器', '操作系统', '状态', '失败原因', '登出时间']
    rows = []
    for log in logs:
        rows.append([
            log.login_time.strftime('%Y-%m-%d %H:%M:%S') if log.login_time else '',
            log.username or '',
            log.ip_address or '',
            log.browser or '',
            log.os or '',
            '成功' if log.status == 'success' else '失败',
            log.fail_reason or '',
            log.logout_time.strftime('%Y-%m-%d %H:%M:%S') if log.logout_time else ''
        ])
    data = export_to_excel(headers, rows, '登录日志', 'login_logs.xlsx')
    return send_file(data, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'登录日志_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx')


# ============== 错误日志 ==============

@bp.route('/error_logs')
@login_required
@admin_required
def error_logs():
    from app.models import ErrorLog
    page = request.args.get('page', 1, type=int)
    error_type = request.args.get('error_type', '', type=str)
    is_handled = request.args.get('is_handled', '', type=str)

    query = ErrorLog.query
    if error_type:
        query = query.filter(ErrorLog.error_type.contains(error_type))
    if is_handled:
        query = query.filter_by(is_handled=(is_handled == '1'))

    pagination = query.order_by(ErrorLog.error_time.desc()).paginate(page=page, per_page=20, error_out=False)
    return render_template('admin/error_logs.html', pagination=pagination, error_type=error_type, is_handled=is_handled)


@bp.route('/error_logs/<int:id>')
@login_required
@admin_required
def error_log_detail(id):
    from app.models import ErrorLog
    log = ErrorLog.query.get_or_404(id)
    return render_template('admin/error_log_detail.html', log=log)


@bp.route('/error_logs/<int:id>/handle', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='处理错误日志')
def handle_error_log(id):
    from app.models import ErrorLog
    log = ErrorLog.query.get_or_404(id)
    log.is_handled = True
    log.handled_by = current_user.name or current_user.username
    log.handled_at = datetime.now()
    log.handle_note = request.form.get('note', '')
    db.session.commit()
    flash('错误日志已标记为已处理', 'success')
    return redirect(url_for('admin.error_log_detail', id=id))


@bp.route('/error_logs/cleanup', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='清理错误日志')
def cleanup_error_logs():
    from app.models import ErrorLog
    days = int(request.form.get('days', 30))
    from datetime import timedelta
    cutoff = datetime.now() - timedelta(days=days)
    count = ErrorLog.query.filter(ErrorLog.error_time < cutoff).delete()
    db.session.commit()
    flash(f'已清理 {count} 条 {days} 天前的错误日志', 'success')
    return redirect(url_for('admin.error_logs'))


# ============== 单据编号检查 ==============

@bp.route('/code_check')
@login_required
@admin_required
def code_check():
    from app.models import StockIn, StockOut, Contract, Payment, Reconciliation
    project_id = session.get('current_project_id')
    if not project_id:
        flash('请先选择项目', 'warning')
        return redirect(url_for('main.index'))

    doc_type = request.args.get('doc_type', 'stock_in', type=str)
    start_date = request.args.get('start_date', '', type=str)
    end_date = request.args.get('end_date', '', type=str)

    result = None
    if request.args.get('check'):
        doc_models = {
            'stock_in': (StockIn, 'code', 'stock_in_date', '入库单'),
            'stock_out': (StockOut, 'code', 'stock_out_date', '出库单'),
            'contract': (Contract, 'code', 'sign_date', '合同'),
            'payment': (Payment, 'code', 'payment_date', '付款单'),
            'reconciliation': (Reconciliation, 'code', 'reconcile_date', '对账单'),
        }
        model, code_field, date_field, label = doc_models.get(doc_type, (StockIn, 'code', 'stock_in_date', '入库单'))

        query = model.query.filter(model.project_id == project_id)
        if start_date:
            query = query.filter(getattr(model, date_field) >= datetime.strptime(start_date, '%Y-%m-%d').date())
        if end_date:
            query = query.filter(getattr(model, date_field) <= datetime.strptime(end_date, '%Y-%m-%d').date())

        docs = query.order_by(getattr(model, code_field)).all()
        codes = [getattr(d, code_field) for d in docs]

        missing = []
        if len(codes) >= 2:
            prefix_parts = codes[0].rsplit('-', 1)
            if len(prefix_parts) == 2 and prefix_parts[1].isdigit():
                prefix = prefix_parts[0] + '-'
                nums = []
                for c in codes:
                    parts = c.rsplit('-', 1)
                    if len(parts) == 2 and parts[1].isdigit():
                        nums.append(int(parts[1]))
                nums.sort()
                if nums:
                    for i in range(nums[0], nums[-1] + 1):
                        if i not in nums:
                            missing.append(prefix + str(i).zfill(len(prefix_parts[1])))

        result = {
            'label': label,
            'total': len(codes),
            'missing_count': len(missing),
            'missing': missing,
            'min_code': codes[0] if codes else '',
            'max_code': codes[-1] if codes else '',
        }

    return render_template('admin/code_check.html', doc_type=doc_type, start_date=start_date, end_date=end_date, result=result)


# ============== 附件管理 ==============

@bp.route('/attachments')
@login_required
@admin_required
def attachments():
    from app.models import Attachment
    page = request.args.get('page', 1, type=int)
    module = request.args.get('module', '', type=str)

    query = Attachment.query.filter_by(is_deleted=False)
    if module:
        query = query.filter_by(module=module)

    pagination = query.order_by(Attachment.uploaded_at.desc()).paginate(page=page, per_page=20, error_out=False)

    total_size = db.session.query(db.func.sum(Attachment.file_size)).filter_by(is_deleted=False).scalar() or 0
    module_stats = db.session.query(
        Attachment.module, db.func.sum(Attachment.file_size)
    ).filter_by(is_deleted=False).group_by(Attachment.module).all()

    return render_template('admin/attachments.html', pagination=pagination, module=module,
                           total_size=total_size, module_stats=module_stats)


@bp.route('/attachments/<int:id>/delete', methods=['POST'])
@login_required
@admin_required
@log_audit(module='admin', operation='删除附件')
def delete_attachment(id):
    from app.models import Attachment
    att = Attachment.query.get_or_404(id)
    att.is_deleted = True
    att.deleted_at = datetime.now()
    db.session.commit()
    flash('附件已删除', 'success')
    return redirect(url_for('admin.attachments'))


@bp.route('/attachments/<int:id>/download')
@login_required
def download_attachment(id):
    """附件下载（带鉴权）"""
    from app.models import Attachment
    from flask import send_from_directory, current_app, abort, session
    import os
    att = Attachment.query.get_or_404(id)
    if att.is_deleted:
        abort(404)
    # 简单鉴权：管理员或同一项目用户可下载
    project_id = session.get('current_project_id')
    if not current_user.is_admin() and att.project_id and att.project_id != project_id:
        abort(403)
    file_dir = os.path.join(current_app.config['UPLOAD_FOLDER'], att.module)
    return send_from_directory(file_dir, att.file_name, as_attachment=True, download_name=att.original_name)


@bp.route('/announcements')
@login_required
@admin_required
def announcements():
    """公告管理列表"""
    page = request.args.get('page', 1, type=int)
    per_page = 20
    from app.models import SysAnnouncement
    query = SysAnnouncement.query.order_by(SysAnnouncement.created_at.desc())
    pagination = query.paginate(page=page, per_page=per_page, error_out=False)
    return render_template('admin/announcements.html', pagination=pagination)


@bp.route('/announcements/create', methods=['GET', 'POST'])
@login_required
@admin_required
@log_audit(module='系统管理', operation='新增公告')
def create_announcement():
    if request.method == 'POST':
        from app.models import SysAnnouncement
        from datetime import datetime
        title = request.form.get('title', '').strip()
        content = request.form.get('content', '').strip()
        type = request.form.get('type', 'notice')
        is_popup = request.form.get('is_popup') == '1'
        status = request.form.get('status') == '1'
        publish_time_str = request.form.get('publish_time', '').strip()
        expire_time_str = request.form.get('expire_time', '').strip()
        visible_scope = request.form.get('visible_scope', 'all')
        visible_roles = request.form.getlist('visible_roles')
        visible_depts = request.form.getlist('visible_depts')

        publish_time = None
        if publish_time_str:
            try:
                publish_time = datetime.strptime(publish_time_str, '%Y-%m-%d %H:%M:%S')
            except ValueError:
                try:
                    publish_time = datetime.strptime(publish_time_str, '%Y-%m-%dT%H:%M')
                except ValueError:
                    publish_time = datetime.now()

        expire_time = None
        if expire_time_str:
            try:
                expire_time = datetime.strptime(expire_time_str, '%Y-%m-%d %H:%M:%S')
            except ValueError:
                try:
                    expire_time = datetime.strptime(expire_time_str, '%Y-%m-%dT%H:%M')
                except ValueError:
                    expire_time = None

        import json
        ann = SysAnnouncement(
            title=title,
            content=content,
            type=type,
            is_popup=is_popup,
            status=status,
            publish_time=publish_time or datetime.now(),
            expire_time=expire_time,
            created_by=current_user.username,
            created_by_id=current_user.id,
            visible_scope=visible_scope,
            visible_roles=json.dumps(visible_roles) if visible_roles else None,
            visible_depts=json.dumps(visible_depts) if visible_depts else None
        )
        db.session.add(ann)
        db.session.commit()
        flash('公告创建成功', 'success')
        return redirect(url_for('admin.announcements'))

    from app.models import SysRole, SysDept
    roles = SysRole.query.all()
    depts = SysDept.query.all()
    return render_template('admin/announcement_form.html', announcement=None, roles=roles, depts=depts)


@bp.route('/announcements/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@admin_required
@log_audit(module='系统管理', operation='编辑公告')
def edit_announcement(id):
    from app.models import SysAnnouncement
    ann = SysAnnouncement.query.get_or_404(id)
    if request.method == 'POST':
        from datetime import datetime
        ann.title = request.form.get('title', '').strip()
        ann.content = request.form.get('content', '').strip()
        ann.type = request.form.get('type', 'notice')
        ann.is_popup = request.form.get('is_popup') == '1'
        ann.status = request.form.get('status') == '1'
        ann.visible_scope = request.form.get('visible_scope', 'all')
        visible_roles = request.form.getlist('visible_roles')
        visible_depts = request.form.getlist('visible_depts')
        import json
        ann.visible_roles = json.dumps(visible_roles) if visible_roles else None
        ann.visible_depts = json.dumps(visible_depts) if visible_depts else None

        publish_time_str = request.form.get('publish_time', '').strip()
        if publish_time_str:
            try:
                ann.publish_time = datetime.strptime(publish_time_str, '%Y-%m-%d %H:%M:%S')
            except ValueError:
                try:
                    ann.publish_time = datetime.strptime(publish_time_str, '%Y-%m-%dT%H:%M')
                except ValueError:
                    pass

        expire_time_str = request.form.get('expire_time', '').strip()
        if expire_time_str:
            try:
                ann.expire_time = datetime.strptime(expire_time_str, '%Y-%m-%d %H:%M:%S')
            except ValueError:
                try:
                    ann.expire_time = datetime.strptime(expire_time_str, '%Y-%m-%dT%H:%M')
                except ValueError:
                    ann.expire_time = None
        else:
            ann.expire_time = None

        db.session.commit()
        flash('公告更新成功', 'success')
        return redirect(url_for('admin.announcements'))

    from app.models import SysRole, SysDept
    roles = SysRole.query.all()
    depts = SysDept.query.all()
    return render_template('admin/announcement_form.html', announcement=ann, roles=roles, depts=depts)


@bp.route('/announcements/<int:id>/delete', methods=['POST'])
@login_required
@admin_required
@log_audit(module='系统管理', operation='删除公告')
def delete_announcement(id):
    from app.models import SysAnnouncement, SysAnnouncementRead
    ann = SysAnnouncement.query.get_or_404(id)
    SysAnnouncementRead.query.filter_by(announcement_id=id).delete()
    db.session.delete(ann)
    db.session.commit()
    flash('公告已删除', 'success')
    return redirect(url_for('admin.announcements'))


@bp.route('/online_users')
@login_required
@admin_required
def online_users():
    """在线用户列表"""
    from app.models import LoginLog, User, SysDept, SysRole
    from datetime import datetime, timedelta

    minutes = request.args.get('minutes', 30, type=int)
    if minutes < 1:
        minutes = 30
    cutoff = datetime.now() - timedelta(minutes=minutes)

    # 查询最近N分钟内活跃的登录记录
    query = LoginLog.query.filter(
        LoginLog.status == 'success',
        LoginLog.user_id != None,
        LoginLog.logout_time == None,
        db.or_(LoginLog.last_active_at >= cutoff, LoginLog.login_time >= cutoff)
    ).order_by(LoginLog.last_active_at.desc().nullslast(), LoginLog.login_time.desc())

    all_logins = query.all()

    # 每个用户只保留最新一条
    seen_users = set()
    online_logins = []
    for log in all_logins:
        if log.user_id not in seen_users:
            seen_users.add(log.user_id)
            online_logins.append(log)

    # 补充用户信息
    user_ids = [l.user_id for l in online_logins if l.user_id]
    users = {u.id: u for u in User.query.filter(User.id.in_(user_ids)).all()} if user_ids else {}
    depts = {d.id: d for d in SysDept.query.all()}
    roles = {r.id: r for r in SysRole.query.all()}

    return render_template('admin/online_users.html',
                           online_logins=online_logins,
                           users=users, depts=depts, roles=roles,
                           minutes=minutes)


@bp.route('/online_users/<int:login_id>/kick', methods=['POST'])
@login_required
@admin_required
@log_audit(module='系统管理', operation='强制下线')
def kick_user(login_id):
    """强制下线用户"""
    from app.models import LoginLog
    from datetime import datetime
    login_log = LoginLog.query.get_or_404(login_id)
    login_log.logout_time = datetime.now()
    db.session.commit()
    flash('已强制下线该用户', 'success')
    return redirect(url_for('admin.online_users'))


@bp.route('/db_migrations')
@login_required
@admin_required
def db_migrations():
    """数据库版本管理"""
    from app.migration import get_applied_versions, get_pending_migrations
    applied = get_applied_versions()
    pending = get_pending_migrations()
    return render_template('admin/db_migrations.html',
                           applied=applied, pending=pending,
                           current_version=applied[-1]['version'] if applied else '0')


@bp.route('/db_migrations/run', methods=['POST'])
@login_required
@admin_required
@log_audit(module='系统管理', operation='执行数据库迁移')
def run_db_migrations():
    """执行待处理的迁移"""
    from app.migration import run_migrations
    results = run_migrations()
    success_count = sum(1 for r in results if r['success'])
    fail_count = sum(1 for r in results if not r['success'])
    if fail_count > 0:
        flash(f'迁移完成：成功 {success_count} 个，失败 {fail_count} 个', 'warning')
    else:
        flash(f'迁移完成：成功执行 {success_count} 个迁移', 'success')
    return redirect(url_for('admin.db_migrations'))
