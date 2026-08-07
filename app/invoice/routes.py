# -*- coding: utf-8 -*-
"""M5 票据三流合一 —— 路由"""

from flask import (render_template, request, redirect, url_for, flash,
                   jsonify, current_app, session)
from flask_login import login_required, current_user

from app import db
from app.invoice import invoice_bp
from app.invoice.models import InvoiceThreeFlowCheck
from app.invoice.services import refresh_three_flow, book_three_flow
from app.models import (Invoice, Payment, Reconciliation, Supplier, Contract,
                        Project)
from app.utils import apply_data_scope
from app.decorators import permission_required


def _current_pid():
    return session.get('current_project_id')


def _fmt(v):
    try:
        return float(v or 0)
    except Exception:
        return 0.0


@invoice_bp.route('/three-flow/')
@login_required
@permission_required('invoice:threeflow:view')
def three_flow_list():
    pid = _current_pid()
    q = InvoiceThreeFlowCheck.query
    if pid:
        q = q.filter_by(project_id=pid)
    q = apply_data_scope(q, InvoiceThreeFlowCheck)

    level = request.args.get('level', '').strip()
    sid = request.args.get('supplier_id', '').strip()
    if level in ('green', 'yellow', 'red'):
        q = q.filter(InvoiceThreeFlowCheck.consistency_level == level)
    if sid:
        q = q.filter(InvoiceThreeFlowCheck.supplier_id == int(sid))

    rows = q.order_by(InvoiceThreeFlowCheck.consistency_level != 'red',
                      InvoiceThreeFlowCheck.id.desc()).all()

    # 汇总
    total = len(rows)
    cnt = {'green': 0, 'yellow': 0, 'red': 0}
    for r in rows:
        cnt[r.consistency_level] = cnt.get(r.consistency_level, 0) + 1

    suppliers = Supplier.query.filter_by(project_id=pid).all() if pid else Supplier.query.all()
    return render_template('invoice/three_flow_list.html',
                           rows=rows, cnt=cnt, total=total,
                           level=level, suppliers=suppliers,
                           Supplier=Supplier, Contract=Contract, Project=Project)


@invoice_bp.route('/three-flow/<int:check_id>/')
@login_required
@permission_required('invoice:threeflow:view')
def three_flow_detail(check_id):
    pid = _current_pid()
    chk = InvoiceThreeFlowCheck.query.get_or_404(check_id)
    if pid and chk.project_id != pid:
        flash('无权限查看其他项目数据', 'danger')
        return redirect(url_for('invoice.three_flow_list'))

    invoices = Invoice.query.filter_by(
        project_id=chk.project_id, supplier_id=chk.supplier_id,
        contract_id=chk.contract_id).order_by(Invoice.invoice_date.desc()).all()
    payments = Payment.query.filter_by(
        project_id=chk.project_id, supplier_id=chk.supplier_id,
        contract_id=chk.contract_id, approval_status='passed').order_by(
        Payment.payment_date.desc()).all()
    reconciliations = Reconciliation.query.filter_by(
        project_id=chk.project_id, supplier_id=chk.supplier_id,
        contract_id=chk.contract_id, status='approved').order_by(
        Reconciliation.end_date.desc()).all()

    supplier = Supplier.query.get(chk.supplier_id)
    contract = Contract.query.get(chk.contract_id)
    return render_template('invoice/three_flow_detail.html',
                           chk=chk, invoices=invoices, payments=payments,
                           reconciliations=reconciliations,
                           supplier=supplier, contract=contract,
                           project=Project.query.get(chk.project_id),
                           fmt=_fmt)


@invoice_bp.route('/three-flow/refresh', methods=['POST'])
@login_required
@permission_required('invoice:threeflow:view')
def three_flow_refresh():
    pid = _current_pid()
    if not pid:
        flash('请先选择项目', 'warning')
        return redirect(url_for('invoice.three_flow_list'))
    try:
        n = refresh_three_flow(pid, operator=getattr(current_user, 'username', None))
        flash('已重算 %d 个供应商×合同核对单元' % n, 'success')
    except Exception as e:
        current_app.logger.warning('[M5] refresh failed: %s' % e)
        flash('重算失败：%s' % e, 'danger')
    return redirect(url_for('invoice.three_flow_list'))


@invoice_bp.route('/three-flow/<int:check_id>/book', methods=['POST'])
@login_required
@permission_required('invoice:threeflow:book')
def three_flow_book(check_id):
    chk = InvoiceThreeFlowCheck.query.get_or_404(check_id)
    ok, msg = book_three_flow(
        check_id, operator=getattr(current_user, 'username', None))
    flash(msg, 'success' if ok else 'danger')
    return redirect(url_for('invoice.three_flow_detail', check_id=check_id))


@invoice_bp.route('/three-flow/export')
@login_required
@permission_required('invoice:threeflow:view')
def three_flow_export():
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    pid = _current_pid()
    q = InvoiceThreeFlowCheck.query
    if pid:
        q = q.filter_by(project_id=pid)
    q = apply_data_scope(q, InvoiceThreeFlowCheck)
    rows = q.order_by(InvoiceThreeFlowCheck.id.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = '三流合一核对'
    headers = ['供应商', '合同', '项目', '票流(含税)', '资金流(含税)',
               '货流(含税)', '合同额(含税)', '票-资差', '资-货差', '票-货差',
               '一致性', '不一致类型', '已入账']
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)

    for r in rows:
        supplier = Supplier.query.get(r.supplier_id)
        contract = Contract.query.get(r.contract_id)
        project = Project.query.get(r.project_id)
        ws.append([
            supplier.name if supplier else r.supplier_id,
            contract.name if contract else r.contract_id,
            project.name if project else r.project_id,
            _fmt(r.invoice_total), _fmt(r.payment_total), _fmt(r.goods_total),
            _fmt(r.contract_amount), _fmt(r.variance_inv_pay),
            _fmt(r.variance_pay_goods), _fmt(r.variance_inv_goods),
            {'green': '一致', 'yellow': '需关注', 'red': '不一致'}.get(
                r.consistency_level, r.consistency_level),
            r.mismatch_types or '-',
            '是' if r.booked else '否',
        ])
    # 颜色标注一致性列（第11列）
    for ri in range(2, len(rows) + 2):
        cell = ws.cell(row=ri, column=11)
        lvl = cell.value
        if lvl == '一致':
            cell.fill = PatternFill('solid', fgColor='C6EFCE')
        elif lvl == '需关注':
            cell.fill = PatternFill('solid', fgColor='FFEB9C')
        else:
            cell.fill = PatternFill('solid', fgColor='FFC7CE')

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    from flask import Response
    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment;filename=three_flow_%s.xlsx' %
                 (pid or 'all')})
