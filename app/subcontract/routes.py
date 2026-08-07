from flask import (render_template, request, redirect, url_for, flash,
                   jsonify, session, make_response)
from flask_login import login_required, current_user
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from app import db
from app.models import (SubcontractDeduction, UsageUnit, StockOut, StockOutItem,
                        Material, Category, Project)
from datetime import datetime, date
from sqlalchemy import func, extract

from app.subcontract import bp
from app.cost.services import safe_record_spend


def _get_period_filters(period):
    """根据 YYYY-MM 解析月份起止日期"""
    try:
        year, month = map(int, period.split('-'))
        start = date(year, month, 1)
        if month == 12:
            end = date(year + 1, 1, 1)
        else:
            end = date(year, month + 1, 1)
        return start, end
    except (ValueError, AttributeError):
        return None, None


@bp.route('/')
@login_required
def index():
    """扣款台账列表"""
    project_id = session.get('current_project_id')
    if not project_id:
        flash('请先选择项目。', 'warning')
        return redirect(url_for('main.index'))

    period = request.args.get('period', '', type=str).strip()
    usage_unit_id = request.args.get('usage_unit_id', 0, type=int)
    status = request.args.get('status', '', type=str).strip()

    query = SubcontractDeduction.query.filter_by(project_id=project_id)
    if period:
        query = query.filter(SubcontractDeduction.period == period)
    if usage_unit_id:
        query = query.filter(SubcontractDeduction.usage_unit_id == usage_unit_id)
    if status:
        query = query.filter(SubcontractDeduction.status == status)

    deductions = query.order_by(
        SubcontractDeduction.period.desc(),
        SubcontractDeduction.created_at.desc()
    ).all()

    units = UsageUnit.query.filter_by(project_id=project_id, is_subcontractor=True) \
        .order_by(UsageUnit.name.asc()).all()

    return render_template('subcontract/list.html', deductions=deductions,
                           units=units, period=period,
                           usage_unit_id=usage_unit_id, status=status)


@bp.route('/generate', methods=['GET', 'POST'])
@login_required
def generate():
    """生成月度汇总"""
    project_id = session.get('current_project_id')
    if not project_id:
        flash('请先选择项目。', 'warning')
        return redirect(url_for('main.index'))

    if request.method == 'POST':
        period = request.form.get('period', '').strip()
        if not period:
            flash('请选择月份。', 'danger')
            return redirect(url_for('subcontract.generate'))

        start, end = _get_period_filters(period)
        if not start:
            flash('月份格式错误，应为 YYYY-MM。', 'danger')
            return redirect(url_for('subcontract.generate'))

        # 查询该月所有分包单位的出库记录
        subcontractor_ids = db.session.query(UsageUnit.id).filter(
            UsageUnit.project_id == project_id,
            UsageUnit.is_subcontractor.is_(True)
        ).all()
        subcontractor_id_list = [u[0] for u in subcontractor_ids]
        if not subcontractor_id_list:
            flash('当前项目暂无分包单位，请先在分包单位管理中标记。', 'warning')
            return redirect(url_for('subcontract.units'))

        # 按用料单位汇总出库明细
        summary = db.session.query(
            StockOut.usage_unit_id,
            func.count(StockOut.id).label('order_count'),
            func.sum(StockOutItem.quantity).label('total_qty'),
            func.sum(StockOutItem.amount).label('total_amt')
        ).join(
            StockOutItem, StockOutItem.stock_out_id == StockOut.id
        ).filter(
            StockOut.project_id == project_id,
            StockOut.usage_unit_id.in_(subcontractor_id_list),
            StockOut.stock_out_date >= start,
            StockOut.stock_out_date < end
        ).group_by(StockOut.usage_unit_id).all()

        if not summary:
            flash(f'{period} 月份无可汇总的分包领料记录。', 'info')
            return redirect(url_for('subcontract.generate'))

        created_count = 0
        updated_count = 0
        for row in summary:
            unit_id = row.usage_unit_id
            total_qty = float(row.total_qty or 0)
            total_amt = float(row.total_amt or 0)
            deduction = SubcontractDeduction.query.filter_by(
                project_id=project_id,
                usage_unit_id=unit_id,
                period=period
            ).first()
            if deduction:
                deduction.total_quantity = total_qty
                deduction.total_amount = total_amt
                deduction.updated_at = datetime.now()
                updated_count += 1
            else:
                deduction = SubcontractDeduction(
                    project_id=project_id,
                    usage_unit_id=unit_id,
                    period=period,
                    total_quantity=total_qty,
                    total_amount=total_amt,
                    status='draft'
                )
                db.session.add(deduction)
                created_count += 1
        db.session.commit()

        flash(f'汇总完成：新增 {created_count} 条，更新 {updated_count} 条。', 'success')
        return redirect(url_for('subcontract.index', period=period))

    # GET：预览生成结果
    period = request.args.get('period', datetime.now().strftime('%Y-%m'), type=str).strip()
    start, end = _get_period_filters(period)

    preview = []
    if start:
        subcontractor_ids = db.session.query(UsageUnit.id).filter(
            UsageUnit.project_id == project_id,
            UsageUnit.is_subcontractor.is_(True)
        ).all()
        subcontractor_id_list = [u[0] for u in subcontractor_ids]

        if subcontractor_id_list:
            summary = db.session.query(
                StockOut.usage_unit_id,
                UsageUnit.name.label('unit_name'),
                UsageUnit.code.label('unit_code'),
                func.count(StockOut.id).label('order_count'),
                func.sum(StockOutItem.quantity).label('total_qty'),
                func.sum(StockOutItem.amount).label('total_amt')
            ).join(
                StockOutItem, StockOutItem.stock_out_id == StockOut.id
            ).join(
                UsageUnit, UsageUnit.id == StockOut.usage_unit_id
            ).filter(
                StockOut.project_id == project_id,
                StockOut.usage_unit_id.in_(subcontractor_id_list),
                StockOut.stock_out_date >= start,
                StockOut.stock_out_date < end
            ).group_by(StockOut.usage_unit_id).all()

            for row in summary:
                # 检查是否已存在
                existing = SubcontractDeduction.query.filter_by(
                    project_id=project_id,
                    usage_unit_id=row.usage_unit_id,
                    period=period
                ).first()
                preview.append({
                    'usage_unit_id': row.usage_unit_id,
                    'unit_name': row.unit_name,
                    'unit_code': row.unit_code,
                    'order_count': int(row.order_count or 0),
                    'total_qty': float(row.total_qty or 0),
                    'total_amt': float(row.total_amt or 0),
                    'existing': existing is not None,
                    'existing_status': existing.status if existing else None,
                })

    return render_template('subcontract/generate.html', period=period, preview=preview)


@bp.route('/<int:id>/detail')
@login_required
def detail(id):
    """扣款明细"""
    deduction = SubcontractDeduction.query.get_or_404(id)

    start, end = _get_period_filters(deduction.period)

    # 该月该用料单位的所有出库单及明细
    stock_outs = []
    if start:
        stock_outs = StockOut.query.filter(
            StockOut.project_id == deduction.project_id,
            StockOut.usage_unit_id == deduction.usage_unit_id,
            StockOut.stock_out_date >= start,
            StockOut.stock_out_date < end
        ).order_by(StockOut.stock_out_date.asc(), StockOut.code.asc()).all()

    # 按物资分类汇总
    category_summary = []
    if stock_outs:
        so_ids = [so.id for so in stock_outs]
        rows = db.session.query(
            Category.id,
            Category.name,
            func.sum(StockOutItem.quantity).label('qty'),
            func.sum(StockOutItem.amount).label('amt')
        ).join(
            Material, Material.category_id == Category.id
        ).join(
            StockOutItem, StockOutItem.material_id == Material.id
        ).filter(
            StockOutItem.stock_out_id.in_(so_ids)
        ).group_by(Category.id, Category.name).order_by(Category.id.asc()).all()

        for row in rows:
            category_summary.append({
                'category_id': row.id,
                'category_name': row.name,
                'qty': float(row.qty or 0),
                'amt': float(row.amt or 0),
            })

    # 物资明细汇总
    material_details = []
    if stock_outs:
        so_ids = [so.id for so in stock_outs]
        rows = db.session.query(
            Material.id,
            Material.name,
            Material.code,
            Material.specification,
            Material.unit,
            func.sum(StockOutItem.quantity).label('qty'),
            func.sum(StockOutItem.amount).label('amt')
        ).join(
            StockOutItem, StockOutItem.material_id == Material.id
        ).filter(
            StockOutItem.stock_out_id.in_(so_ids)
        ).group_by(Material.id).order_by(Material.name.asc()).all()

        for row in rows:
            material_details.append({
                'material_id': row.id,
                'name': row.name,
                'code': row.code,
                'spec': row.specification,
                'unit': row.unit,
                'qty': float(row.qty or 0),
                'amt': float(row.amt or 0),
            })

    return render_template('subcontract/detail.html', deduction=deduction,
                           stock_outs=stock_outs,
                           category_summary=category_summary,
                           material_details=material_details)


@bp.route('/<int:id>/export')
@login_required
def export(id):
    """导出扣款单（Excel）"""
    deduction = SubcontractDeduction.query.get_or_404(id)
    project = Project.query.get(deduction.project_id)
    unit = UsageUnit.query.get(deduction.usage_unit_id)

    start, end = _get_period_filters(deduction.period)

    # 收集出库明细
    items = []
    if start:
        so_ids = [so.id for so in StockOut.query.filter(
            StockOut.project_id == deduction.project_id,
            StockOut.usage_unit_id == deduction.usage_unit_id,
            StockOut.stock_out_date >= start,
            StockOut.stock_out_date < end
        ).all()]
        if so_ids:
            items = db.session.query(
                StockOut.code.label('so_code'),
                StockOut.stock_out_date.label('so_date'),
                Material.name.label('material_name'),
                Material.code.label('material_code'),
                Material.specification.label('spec'),
                Material.unit.label('unit'),
                Category.name.label('category_name'),
                StockOutItem.quantity.label('qty'),
                StockOutItem.unit_price.label('price'),
                StockOutItem.amount.label('amt')
            ).join(
                StockOutItem, StockOutItem.stock_out_id == StockOut.id
            ).join(
                Material, Material.id == StockOutItem.material_id
            ).outerjoin(
                Category, Category.id == Material.category_id
            ).filter(
                StockOutItem.stock_out_id.in_(so_ids)
            ).order_by(StockOut.stock_out_date.asc(), StockOut.code.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = '分包扣款单'

    title_font = Font(name='宋体', size=16, bold=True)
    header_font = Font(name='宋体', size=10, bold=True)
    cell_font = Font(name='宋体', size=10)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right = Alignment(horizontal='right', vertical='center')
    thin = Side(border_style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    # 标题行
    title_text = f'{project.name if project else ""}{deduction.period}月份分包扣款单'
    ws.cell(row=1, column=1, value=title_text).font = title_font
    ws.cell(row=1, column=1).alignment = center
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws.row_dimensions[1].height = 30

    # 用料单位信息
    ws.cell(row=2, column=1, value='分包单位：').font = cell_font
    ws.cell(row=2, column=1).alignment = right
    ws.cell(row=2, column=2, value=unit.name if unit else '').font = cell_font
    ws.cell(row=2, column=2).alignment = left
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=4)

    ws.cell(row=2, column=5, value='单位编码：').font = cell_font
    ws.cell(row=2, column=5).alignment = right
    ws.cell(row=2, column=6, value=unit.code or '' if unit else '').font = cell_font
    ws.cell(row=2, column=6).alignment = left
    ws.merge_cells(start_row=2, start_column=6, end_row=2, end_column=8)

    ws.cell(row=3, column=1, value='分包合同：').font = cell_font
    ws.cell(row=3, column=1).alignment = right
    ws.cell(row=3, column=2, value=(unit.subcontract_contract or '') if unit else '').font = cell_font
    ws.cell(row=3, column=2).alignment = left
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=4)

    ws.cell(row=3, column=5, value='月份：').font = cell_font
    ws.cell(row=3, column=5).alignment = right
    ws.cell(row=3, column=6, value=deduction.period).font = cell_font
    ws.cell(row=3, column=6).alignment = left
    ws.merge_cells(start_row=3, start_column=6, end_row=3, end_column=8)

    # 明细表头
    header_row = 5
    headers = ['出库单号', '出库日期', '物资分类', '物资编码', '物资名称', '规格型号', '单位', '数量', '单价', '金额']
    for col_idx, header in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=header)
        c.font = header_font
        c.alignment = center
        c.border = border
        c.fill = header_fill
    ws.row_dimensions[header_row].height = 22

    # 明细数据
    row_idx = header_row + 1
    for item in items:
        row_data = [
            item.so_code or '',
            item.so_date.strftime('%Y-%m-%d') if item.so_date else '',
            item.category_name or '',
            item.material_code or '',
            item.material_name or '',
            item.spec or '',
            item.unit or '',
            float(item.qty or 0),
            float(item.price or 0),
            float(item.amt or 0),
        ]
        for col_idx, val in enumerate(row_data, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font = cell_font
            c.border = border
            if col_idx in (8, 9, 10):
                c.alignment = right
            elif col_idx in (1, 2):
                c.alignment = center
            else:
                c.alignment = left
        row_idx += 1

    # 合计行
    total_qty_cell = ws.cell(row=row_idx, column=7, value='合计')
    total_qty_cell.font = header_font
    total_qty_cell.alignment = right
    total_qty_cell.fill = header_fill
    total_qty_cell.border = border
    qty_cell = ws.cell(row=row_idx, column=8, value=float(deduction.total_quantity or 0))
    qty_cell.font = header_font
    qty_cell.alignment = right
    qty_cell.fill = header_fill
    qty_cell.border = border
    # 占位空列9
    ws.cell(row=row_idx, column=9, value='').fill = header_fill
    ws.cell(row=row_idx, column=9).border = border
    amt_cell = ws.cell(row=row_idx, column=10, value=float(deduction.total_amount or 0))
    amt_cell.font = header_font
    amt_cell.alignment = right
    amt_cell.fill = header_fill
    amt_cell.border = border
    # 合并前面空白列
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
    for col in range(1, 8):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = header_fill
        cell.border = border

    # 列宽
    col_widths = [16, 12, 14, 14, 22, 18, 8, 12, 12, 14]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 备注行
    note_row = row_idx + 2
    ws.cell(row=note_row, column=1, value=f'备注：{deduction.remark or ""}').font = cell_font
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)

    # 制单人/日期
    sign_row = note_row + 1
    ws.cell(row=sign_row, column=1, value=f'制单日期：{datetime.now().strftime("%Y-%m-%d")}').font = cell_font
    ws.merge_cells(start_row=sign_row, start_column=1, end_row=sign_row, end_column=5)
    ws.cell(row=sign_row, column=6, value='制单人：').font = cell_font
    ws.merge_cells(start_row=sign_row, start_column=6, end_row=sign_row, end_column=10)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    resp = make_response(output.getvalue())
    resp.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    filename = f'subcontract_{deduction.period}_{(unit.code or unit.id) if unit else "unknown"}.xlsx'
    resp.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp


@bp.route('/<int:id>/confirm', methods=['POST'])
@login_required
def confirm(id):
    """确认扣款"""
    deduction = SubcontractDeduction.query.get_or_404(id)
    if deduction.status == 'confirmed':
        flash('该扣款记录已确认，无需重复操作。', 'info')
        return redirect(url_for('subcontract.index'))
    deduction.status = 'confirmed'
    deduction.updated_at = datetime.now()
    # M4 P2：记账到预算管控（subcontract 科目），失败不影响主流程
    safe_record_spend(deduction.project_id, 'subcontract',
                      float(deduction.total_amount or 0), 'subcontract', deduction.id)
    db.session.commit()
    flash(f'{deduction.period} 扣款记录已确认。', 'success')
    return redirect(url_for('subcontract.index'))


@bp.route('/units')
@login_required
def units():
    """分包单位管理"""
    project_id = session.get('current_project_id')
    if not project_id:
        flash('请先选择项目。', 'warning')
        return redirect(url_for('main.index'))

    units_list = UsageUnit.query.filter_by(project_id=project_id) \
        .order_by(UsageUnit.is_subcontractor.desc(), UsageUnit.name.asc()).all()

    return render_template('subcontract/units.html', units=units_list)


@bp.route('/units/<int:unit_id>/toggle', methods=['POST'])
@login_required
def toggle_unit(unit_id):
    """切换分包标记"""
    unit = UsageUnit.query.get_or_404(unit_id)
    unit.is_subcontractor = not bool(unit.is_subcontractor)
    db.session.commit()
    return jsonify({
        'success': True,
        'is_subcontractor': unit.is_subcontractor,
        'message': f'已{"标记" if unit.is_subcontractor else "取消标记"}为分包单位'
    })


@bp.route('/units/<int:unit_id>/contract', methods=['POST'])
@login_required
def set_contract(unit_id):
    """设置分包合同"""
    unit = UsageUnit.query.get_or_404(unit_id)
    subcontract_contract = (request.form.get('subcontract_contract', '') or '').strip()
    if not subcontract_contract:
        try:
            data = request.get_json(silent=True) or {}
            subcontract_contract = (data.get('subcontract_contract', '') or '').strip()
        except Exception:
            subcontract_contract = ''
    unit.subcontract_contract = subcontract_contract or None
    db.session.commit()
    return jsonify({
        'success': True,
        'subcontract_contract': unit.subcontract_contract or '',
        'message': '分包合同已保存'
    })
