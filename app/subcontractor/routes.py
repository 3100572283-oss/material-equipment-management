# app/subcontractor/routes.py
"""M6 分包商核验路由：准入档案 + 四维度明细 + 黑名单核查 + 企查查 + 资料包导出 + 名录维护。"""
from datetime import datetime

from flask import (render_template, request, redirect, url_for, flash,
                   jsonify, current_app, session, send_file)
from flask_login import login_required, current_user
from sqlalchemy import or_

from app import db
from app.subcontractor import subcontractor_bp
from app.subcontractor.models import (SubcontractorProfile, SubcontractorQualification,
                                      SubcontractorSafetyCert, SubcontractorPerformance,
                                      SubcontractorBlacklist, SubcontractorBlacklistCheck,
                                      SubcontractorQccSnapshot, SubcontractorAdmitPackage,
                                      run_blacklist_check, BLACKLIST_LISTS, QccAdapter,
                                      ADMIT_STATUS)
from app.models import Supplier, Project
from app.utils import apply_data_scope
from app.decorators import permission_required


def _current_pid():
    return session.get('current_project_id')


def _operator():
    return getattr(current_user, 'username', None) or getattr(current_user, 'name', None)


# ============================================================ 准入档案
@subcontractor_bp.route('/')
@subcontractor_bp.route('/profile/')
@login_required
@permission_required('subcontractor:profile:view')
def profile_list():
    pid = _current_pid()
    q = SubcontractorProfile.query
    if pid:
        q = q.filter_by(project_id=pid)
    q = apply_data_scope(q, SubcontractorProfile)
    keyword = request.args.get('keyword', '').strip()
    status = request.args.get('status', '').strip()
    if keyword:
        q = q.join(Supplier).filter(
            or_(Supplier.name.contains(keyword), Supplier.credit_code.contains(keyword)))
    if status:
        q = q.filter_by(admit_status=status)
    profiles = q.order_by(SubcontractorProfile.created_at.desc()).all()
    return render_template('subcontractor/profile_list.html',
                           profiles=profiles, keyword=keyword, status=status,
                           ADMIT_STATUS=ADMIT_STATUS)


@subcontractor_bp.route('/profile/create', methods=['GET', 'POST'])
@login_required
@permission_required('subcontractor:profile:edit')
def profile_create():
    pid = _current_pid()
    suppliers = Supplier.query.filter_by(is_deleted=False).order_by(Supplier.name).all()
    if request.method == 'POST':
        supplier_id = request.form.get('supplier_id', type=int)
        if not supplier_id:
            flash('请选择关联供应商。', 'warning')
            return render_template('subcontractor/profile_form.html', suppliers=suppliers,
                                   profile=None)
        # 同一供应商同一项目不重复建档
        exists = SubcontractorProfile.query.filter_by(
            supplier_id=supplier_id, project_id=pid).first() if pid else None
        if exists:
            flash('该供应商在本项目已存在准入档案。', 'warning')
            return redirect(url_for('subcontractor.profile_detail', id=exists.id))
        profile = SubcontractorProfile(
            supplier_id=supplier_id, project_id=pid or 0,
            admit_scope=request.form.get('admit_scope', '').strip(),
            admit_status='pending', created_by=_operator())
        vf = request.form.get('valid_from')
        vt = request.form.get('valid_to')
        if vf:
            profile.valid_from = datetime.strptime(vf, '%Y-%m-%d').date()
        if vt:
            profile.valid_to = datetime.strptime(vt, '%Y-%m-%d').date()
        db.session.add(profile)
        db.session.commit()
        run_blacklist_check(profile)
        flash('准入档案已创建，已自动完成 7 类黑名单比对。', 'success')
        return redirect(url_for('subcontractor.profile_detail', id=profile.id))
    return render_template('subcontractor/profile_form.html', suppliers=suppliers,
                           profile=None)


@subcontractor_bp.route('/profile/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@permission_required('subcontractor:profile:edit')
def profile_edit(id):
    profile = SubcontractorProfile.query.get_or_404(id)
    suppliers = Supplier.query.filter_by(is_deleted=False).order_by(Supplier.name).all()
    if request.method == 'POST':
        profile.admit_scope = request.form.get('admit_scope', '').strip()
        vf = request.form.get('valid_from')
        vt = request.form.get('valid_to')
        if vf:
            profile.valid_from = datetime.strptime(vf, '%Y-%m-%d').date()
        if vt:
            profile.valid_to = datetime.strptime(vt, '%Y-%m-%d').date()
        db.session.commit()
        flash('准入档案已更新。', 'success')
        return redirect(url_for('subcontractor.profile_detail', id=profile.id))
    return render_template('subcontractor/profile_form.html', suppliers=suppliers,
                           profile=profile)


@subcontractor_bp.route('/profile/<int:id>/')
@login_required
@permission_required('subcontractor:profile:view')
def profile_detail(id):
    profile = SubcontractorProfile.query.get_or_404(id)
    blacklist_checks = (SubcontractorBlacklistCheck.query
                        .filter_by(profile_id=profile.id).all())
    qcc = (SubcontractorQccSnapshot.query.filter_by(profile_id=profile.id)
           .order_by(SubcontractorQccSnapshot.checked_at.desc()).first())
    return render_template('subcontractor/profile_detail.html',
                           profile=profile, supplier=profile.supplier,
                           qualifications=profile.qualifications,
                           safety_certs=profile.safety_certs,
                           performances=profile.performances,
                           blacklist_checks=blacklist_checks,
                           qcc=qcc, BLACKLIST_LISTS=BLACKLIST_LISTS,
                           ADMIT_STATUS=ADMIT_STATUS)


@subcontractor_bp.route('/profile/<int:id>/approve', methods=['POST'])
@login_required
@permission_required('subcontractor:profile:review')
def profile_approve(id):
    profile = SubcontractorProfile.query.get_or_404(id)
    if profile.blacklist_hit:
        flash('该分包商命中黑名单，禁止准入通过；请先处理黑名单。', 'danger')
        return redirect(url_for('subcontractor.profile_detail', id=profile.id))
    profile.admit_status = 'approved'
    profile.approved_by = _operator()
    profile.approved_at = datetime.now()
    profile.reject_reason = None
    db.session.commit()
    flash('准入审批通过，已纳入合格分包商名录。', 'success')
    return redirect(url_for('subcontractor.profile_detail', id=profile.id))


@subcontractor_bp.route('/profile/<int:id>/reject', methods=['POST'])
@login_required
@permission_required('subcontractor:profile:review')
def profile_reject(id):
    profile = SubcontractorProfile.query.get_or_404(id)
    reason = request.form.get('reject_reason', '').strip()
    profile.admit_status = 'rejected'
    profile.reject_reason = reason
    profile.approved_at = datetime.now()
    db.session.commit()
    flash('准入已驳回。', 'success')
    return redirect(url_for('subcontractor.profile_detail', id=profile.id))


@subcontractor_bp.route('/profile/<int:id>/recheck', methods=['POST'])
@login_required
@permission_required('subcontractor:profile:edit')
def profile_recheck(id):
    profile = SubcontractorProfile.query.get_or_404(id)
    hit = run_blacklist_check(profile)
    # 企查查核验（Adapter 预留）
    supplier = profile.supplier
    adapter = QccAdapter(api_key=current_app.config.get('QCC_API_KEY'))
    result = adapter.verify(credit_code=supplier.credit_code, name=supplier.name)
    snap = SubcontractorQccSnapshot(profile_id=profile.id, status=result.get('status', 'manual'))
    if result.get('configured') and result.get('raw'):
        snap.raw_json = str(result.get('raw'))
    db.session.add(snap)
    profile.qcc_status = result.get('status', 'manual')
    db.session.commit()
    flash('已重新核查：黑名单%s，企查查状态=%s。'
          % ('命中' if hit else '未命中', result.get('status')), 'info')
    return redirect(url_for('subcontractor.profile_detail', id=profile.id))


@subcontractor_bp.route('/profile/<int:id>/export')
@login_required
@permission_required('subcontractor:profile:view')
def profile_export(id):
    profile = SubcontractorProfile.query.get_or_404(id)
    supplier = profile.supplier
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        flash('服务器未安装 python-docx，无法导出资料包。', 'danger')
        return redirect(url_for('subcontractor.profile_detail', id=profile.id))

    doc = Document()
    doc.add_heading('分包商准入资料包', 0)
    doc.add_paragraph('（按铁建云链 crccep.cn 准入要求模板字段对齐，导出后由人工上传）')

    doc.add_heading('一、基本信息', level=1)
    base = [
        ('公司名称', supplier.name if supplier else '-'),
        ('统一社会信用代码', supplier.credit_code or '-' if supplier else '-'),
        ('法定代表人', supplier.legal_person or '-' if supplier else '-'),
        ('联系人', supplier.contact_person or '-' if supplier else '-'),
        ('联系电话', supplier.phone or '-' if supplier else '-'),
        ('注册地址', supplier.address or '-' if supplier else '-'),
        ('开户银行', supplier.bank_name or '-' if supplier else '-'),
        ('银行账号', supplier.bank_account or '-' if supplier else '-'),
        ('营业执照有效期至', str(supplier.license_expire_date) if supplier and supplier.license_expire_date else '-'),
        ('准入范围', profile.admit_scope or '-'),
        ('准入有效期', '%s ~ %s' % (profile.valid_from or '-', profile.valid_to or '-')),
        ('准入结论', ADMIT_STATUS.get(profile.admit_status, profile.admit_status)),
        ('黑名单核查', '命中' if profile.blacklist_hit else '未命中'),
        ('企查查核验', profile.qcc_status),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Light Grid Accent 1'
    for k, v in base:
        row = t.add_row().cells
        row[0].text = k
        row[1].text = str(v)

    def _block(title, rows, headers):
        doc.add_heading(title, level=1)
        if not rows:
            doc.add_paragraph('（暂无记录）')
            return
        tb = doc.add_table(rows=1, cols=len(headers))
        tb.style = 'Light Grid Accent 1'
        for i, h in enumerate(headers):
            tb.rows[0].cells[i].text = h
        for r in rows:
            cells = tb.add_row().cells
            for i, val in enumerate(r):
                cells[i].text = '' if val is None else str(val)

    _block('二、资质条件',
           [[q.cert_type, q.cert_no, q.issuer, q.valid_to] for q in profile.qualifications],
           ['资质类型', '证书编号', '发证机关', '有效期至'])
    _block('三、安全资格',
           [[s.cert_type, s.cert_no, s.valid_to] for s in profile.safety_certs],
           ['证书类型', '证书编号', '有效期至'])
    _block('四、业绩要求',
           [[p.project_name, p.scale, p.period] for p in profile.performances],
           ['工程名称', '规模/金额', '工期/年份'])

    doc.add_heading('五、黑名单核查结果', level=1)
    if profile.blacklist_checks:
        tb = doc.add_table(rows=1, cols=3)
        tb.style = 'Light Grid Accent 1'
        for i, h in enumerate(['名录', '是否命中', '来源']):
            tb.rows[0].cells[i].text = h
        for c in profile.blacklist_checks:
            cells = tb.add_row().cells
            cells[0].text = c.list_name
            cells[1].text = '命中' if c.hit else '未命中'
            cells[2].text = c.source or '-'
    else:
        doc.add_paragraph('（尚未核查）')

    import io
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    fname = '准入资料包_%s.docx' % (supplier.name if supplier else profile.id)
    rec = SubcontractorAdmitPackage(profile_id=profile.id, file_name=fname,
                                    exported_by=_operator())
    db.session.add(rec)
    db.session.commit()
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


# ============================================================ 四维度明细维护
@subcontractor_bp.route('/profile/<int:id>/qualification/add', methods=['POST'])
@login_required
@permission_required('subcontractor:profile:edit')
def qc_add(id):
    profile = SubcontractorProfile.query.get_or_404(id)
    cert_type = request.form.get('cert_type', '').strip()
    if not cert_type:
        flash('请填写资质类型。', 'warning')
        return redirect(url_for('subcontractor.profile_detail', id=profile.id))
    q = SubcontractorQualification(
        profile_id=profile.id, cert_type=cert_type,
        cert_no=request.form.get('cert_no', '').strip(),
        issuer=request.form.get('issuer', '').strip(),
        attachment=request.form.get('attachment', '').strip() or None)
    vt = request.form.get('valid_to')
    if vt:
        q.valid_to = datetime.strptime(vt, '%Y-%m-%d').date()
    db.session.add(q)
    db.session.commit()
    flash('资质已添加。', 'success')
    return redirect(url_for('subcontractor.profile_detail', id=profile.id))


@subcontractor_bp.route('/qualification/<int:qid>/delete', methods=['POST'])
@login_required
@permission_required('subcontractor:profile:edit')
def qc_delete(qid):
    q = SubcontractorQualification.query.get_or_404(qid)
    pid = q.profile_id
    db.session.delete(q)
    db.session.commit()
    flash('资质已删除。', 'success')
    return redirect(url_for('subcontractor.profile_detail', id=pid))


@subcontractor_bp.route('/profile/<int:id>/safety/add', methods=['POST'])
@login_required
@permission_required('subcontractor:profile:edit')
def safety_add(id):
    profile = SubcontractorProfile.query.get_or_404(id)
    cert_type = request.form.get('cert_type', '').strip()
    if not cert_type:
        flash('请填写证书类型。', 'warning')
        return redirect(url_for('subcontractor.profile_detail', id=profile.id))
    s = SubcontractorSafetyCert(
        profile_id=profile.id, cert_type=cert_type,
        cert_no=request.form.get('cert_no', '').strip(),
        attachment=request.form.get('attachment', '').strip() or None)
    vt = request.form.get('valid_to')
    if vt:
        s.valid_to = datetime.strptime(vt, '%Y-%m-%d').date()
    db.session.add(s)
    db.session.commit()
    flash('安全资格已添加。', 'success')
    return redirect(url_for('subcontractor.profile_detail', id=profile.id))


@subcontractor_bp.route('/safety/<int:sid>/delete', methods=['POST'])
@login_required
@permission_required('subcontractor:profile:edit')
def safety_delete(sid):
    s = SubcontractorSafetyCert.query.get_or_404(sid)
    pid = s.profile_id
    db.session.delete(s)
    db.session.commit()
    flash('安全资格已删除。', 'success')
    return redirect(url_for('subcontractor.profile_detail', id=pid))


@subcontractor_bp.route('/profile/<int:id>/performance/add', methods=['POST'])
@login_required
@permission_required('subcontractor:profile:edit')
def perf_add(id):
    profile = SubcontractorProfile.query.get_or_404(id)
    project_name = request.form.get('project_name', '').strip()
    if not project_name:
        flash('请填写工程名称。', 'warning')
        return redirect(url_for('subcontractor.profile_detail', id=profile.id))
    p = SubcontractorPerformance(
        profile_id=profile.id, project_name=project_name,
        scale=request.form.get('scale', '').strip(),
        period=request.form.get('period', '').strip(),
        proof_attachment=request.form.get('proof_attachment', '').strip() or None)
    db.session.add(p)
    db.session.commit()
    flash('业绩已添加。', 'success')
    return redirect(url_for('subcontractor.profile_detail', id=profile.id))


@subcontractor_bp.route('/performance/<int:pid>/delete', methods=['POST'])
@login_required
@permission_required('subcontractor:profile:edit')
def perf_delete(pid):
    p = SubcontractorPerformance.query.get_or_404(pid)
    prof_id = p.profile_id
    db.session.delete(p)
    db.session.commit()
    flash('业绩已删除。', 'success')
    return redirect(url_for('subcontractor.profile_detail', id=prof_id))


# ============================================================ 黑名单名录维护（company 级）
@subcontractor_bp.route('/blacklist/')
@login_required
@permission_required('subcontractor:blacklist:view')
def blacklist_list():
    entries = (SubcontractorBlacklist.query.filter_by(is_active=True)
               .order_by(SubcontractorBlacklist.list_name).all())
    return render_template('subcontractor/blacklist_list.html',
                           entries=entries, BLACKLIST_LISTS=BLACKLIST_LISTS)


@subcontractor_bp.route('/blacklist/add', methods=['POST'])
@login_required
@permission_required('subcontractor:blacklist:edit')
def blacklist_add():
    list_name = request.form.get('list_name', '').strip()
    if list_name not in BLACKLIST_LISTS:
        flash('非法的黑名单名录类型。', 'warning')
        return redirect(url_for('subcontractor.blacklist_list'))
    e = SubcontractorBlacklist(
        list_name=list_name,
        entity_name=request.form.get('entity_name', '').strip() or None,
        credit_code=request.form.get('credit_code', '').strip() or None,
        reason=request.form.get('reason', '').strip() or None,
        source='manual', created_by=_operator())
    db.session.add(e)
    db.session.commit()
    flash('黑名单记录已新增。', 'success')
    return redirect(url_for('subcontractor.blacklist_list'))


@subcontractor_bp.route('/blacklist/<int:id>/delete', methods=['POST'])
@login_required
@permission_required('subcontractor:blacklist:edit')
def blacklist_delete(id):
    e = SubcontractorBlacklist.query.get_or_404(id)
    e.is_active = False
    db.session.commit()
    flash('黑名单记录已移除。', 'success')
    return redirect(url_for('subcontractor.blacklist_list'))
