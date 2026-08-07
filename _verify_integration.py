#!/usr/bin/env python
"""外部对接中心端到端自检（本地 SQLite，不触碰生产）。"""
import os
import sys
from datetime import date, datetime

os.environ.setdefault('SECRET_KEY', 'dev-secret')
os.environ.setdefault('AUTH_CORE_ENABLED', 'true')
os.environ['DATABASE_URL'] = 'sqlite:///' + os.path.abspath('_verify_integ.db')
if os.path.exists('_verify_integ.db'):
    os.remove('_verify_integ.db')

from app import create_app, db  # noqa: E402

app = create_app()
FAIL = []


def check(name, cond, extra=''):
    print(('  ✅ ' if cond else '  ❌ ') + name + (' — ' + str(extra) if extra else ''))
    if not cond:
        FAIL.append(name)


with app.app_context():
    from app.integration.models import (ExtCredential, ExtEndpoint, ExportTemplate,
                                        ExportTemplateField, ExtQueryCache, ExtQueryLog,
                                        ensure_builtin_integration_data, FIELD_SOURCES)
    from app.integration import qcc_client, services
    from app.models import Project, Supplier
    from app.subcontractor.models import (SubcontractorProfile, SubcontractorQualification,
                                          SubcontractorSafetyCert, SubcontractorPerformance,
                                          run_blacklist_check)

    print('\n[1] 内置数据初始化（幂等）')
    ensure_builtin_integration_data()
    n1 = ExtCredential.query.count()
    ensure_builtin_integration_data()          # 第二次不应重复插入
    n2 = ExtCredential.query.count()
    check('凭据占位创建且幂等', n1 == n2 == 2, 'count=%s' % n2)
    check('企查查接口目录 5 条', ExtEndpoint.query.filter_by(provider='qcc').count() == 5)
    tpl = ExportTemplate.query.filter_by(code='crccep_subcontractor_v1').first()
    check('内置铁建云链模板存在', tpl is not None)
    check('模板 28 个字段', tpl and len(tpl.fields) == 28, len(tpl.fields) if tpl else 0)
    check('字段取值源全部在白名单内',
          all((f.source_expr in FIELD_SOURCES) for f in tpl.fields if f.source_expr))

    print('\n[2] 密钥加密存储与掩码')
    from app.utils import encrypt_secret, decrypt_secret
    cred = ExtCredential.query.filter_by(provider='qcc').first()
    cred.api_key = encrypt_secret('abcdef1234567890KEY')
    cred.api_secret = encrypt_secret('S3cretV4lue')
    db.session.commit()
    check('密文以 ENC: 开头', cred.api_key.startswith('ENC:'))
    check('可正确解密', cred.get_key() == 'abcdef1234567890KEY', cred.get_key())
    check('掩码不泄露完整 key', cred.key_masked == 'abcd****0KEY', cred.key_masked)

    print('\n[3] 未启用时优雅降级（关键：不得阻断主流程）')
    cred.enabled = False
    db.session.commit()
    check('is_configured=False', qcc_client.is_configured() is False)
    r = qcc_client.call('basic', '某某建筑有限公司')
    check('call 返回结构化失败而非抛异常', r.get('status') == 'manual', r.get('message'))

    print('\n[4] 造测试数据 → 核验编排降级路径')
    proj = Project(name='测试项目部', code='TP001')
    db.session.add(proj)
    db.session.flush()
    sup = Supplier(project_id=proj.id, name='中铁测试劳务分包有限公司',
                   credit_code='91110000TEST0001XY', legal_person='张三',
                   contact_person='李四', phone='13800000000',
                   address='北京市海淀区某某路1号', bank_name='中国建设银行北京分行',
                   bank_account='11001000000000000001',
                   license_expire_date=date(2030, 12, 31))
    db.session.add(sup)
    db.session.flush()
    prof = SubcontractorProfile(supplier_id=sup.id, project_id=proj.id,
                                admit_scope='土建劳务分包', admit_status='approved',
                                valid_from=date(2026, 1, 1), valid_to=date(2026, 12, 31),
                                approved_by='王五', approved_at=datetime.now())
    db.session.add(prof)
    db.session.flush()
    db.session.add(SubcontractorQualification(profile_id=prof.id, cert_type='营业执照',
                                              cert_no='BL001', issuer='北京市监局',
                                              valid_to=date(2030, 12, 31)))
    db.session.add(SubcontractorQualification(profile_id=prof.id, cert_type='劳务分包资质',
                                              cert_no='LW002', valid_to=date(2027, 6, 30)))
    db.session.add(SubcontractorSafetyCert(profile_id=prof.id, cert_type='安许证',
                                           cert_no='AQ001', valid_to=date(2028, 3, 1)))
    db.session.add(SubcontractorPerformance(profile_id=prof.id,
                                            project_name='某高铁站房工程', scale='1200万元',
                                            period='2024'))
    db.session.commit()
    run_blacklist_check(prof)

    res = services.verify_subcontractor(prof, operator='tester')
    check('未配置时 configured=False', res['configured'] is False)
    check('降级状态为 manual', res['status'] == 'manual', res['message'])
    check('主流程未抛异常且档案仍可用', SubcontractorProfile.query.get(prof.id) is not None)

    print('\n[5] 模板字段取值（全部 28 列）')
    tpl = ExportTemplate.query.filter_by(code='crccep_subcontractor_v1').first()
    vals = {f.target_field: services.resolve_field(prof, f.source_expr, f.default_value)
            for f in tpl.fields}
    check('企业名称正确', vals['企业名称'] == '中铁测试劳务分包有限公司', vals['企业名称'])
    check('信用代码正确', vals['统一社会信用代码'] == '91110000TEST0001XY')
    check('资质类型合并', vals['资质类型'] == '营业执照、劳务分包资质', vals['资质类型'])
    check('资质数量=2', str(vals['资质证书数量']) == '2', vals['资质证书数量'])
    check('安全资格类型', vals['安全资格类型'] == '安许证')
    check('业绩合并', vals['同类工程业绩'] == '某高铁站房工程')
    check('最早到期证书=2027-06-30', vals['证书最早到期日'] == '2027-06-30', vals['证书最早到期日'])
    check('准入结论中文化', vals['准入审批结论'] == '通过', vals['准入审批结论'])
    check('黑名单结论', vals['黑名单核查结论'] == '未命中')
    check('所属项目', vals['所属项目'] == '测试项目部')
    check('报送单位常量非空', bool(vals['报送单位']), vals['报送单位'])
    check('报送日期=今天', vals['报送日期'] == date.today().strftime('%Y-%m-%d'))
    empty_required = [f.target_field for f in tpl.fields
                      if f.required and not vals.get(f.target_field)]
    check('必填项全部有值', not empty_required, empty_required)

    print('\n[6] 取值源白名单防注入')
    check('非法表达式返回空', services.resolve_field(prof, '__class__.__mro__') == '')
    check('未授权属性返回空', services.resolve_field(prof, 'supplier.bank_account_secret') == '')
    check('空表达式走默认值', services.resolve_field(prof, None, '默认X') == '默认X')

    print('\n[7] 必填校验')
    miss = services.validate_template(prof, tpl)
    check('完整数据无缺失', miss == [], miss)
    sup.phone = None
    db.session.commit()
    miss2 = services.validate_template(prof, tpl)
    check('清空电话后检出缺失', '联系电话' in miss2, miss2)
    sup.phone = '13800000000'
    db.session.commit()

    print('\n[8] xlsx 导出')
    buf = services.export_xlsx([prof], tpl)
    data = buf.getvalue()
    check('产出非空 xlsx', len(data) > 3000, '%s bytes' % len(data))
    check('zip 魔数正确', data[:2] == b'PK')
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    check('表头 28 列', ws.max_column == 28, ws.max_column)
    check('数据 1 行', ws.max_row == 2, ws.max_row)
    check('首列表头正确', ws.cell(1, 1).value == '企业名称', ws.cell(1, 1).value)
    check('首行数据正确', ws.cell(2, 1).value == '中铁测试劳务分包有限公司')
    check('冻结首行', ws.freeze_panes == 'A2')

    print('\n[9] 缓存机制')
    ExtQueryCache.query.delete()
    db.session.commit()
    from app.integration.models import ExtQueryCache as C
    k1 = C.build_key('qcc', 'basic', '  测试公司 ')
    k2 = C.build_key('qcc', 'basic', '测试公司')
    check('缓存 key 去空格后一致', k1 == k2)
    qcc_client._write_cache('basic', '测试公司', {'Status': '200', 'Result': {'Name': 'X'}}, 30)
    got = qcc_client._read_cache('basic', '测试公司')
    check('缓存写入并可读回', got and got.get('Status') == '200', got)
    row = C.query.first()
    row.expires_at = datetime(2020, 1, 1)
    db.session.commit()
    check('过期缓存不再返回', qcc_client._read_cache('basic', '测试公司') is None)

    print('\n[10] 启用后真实调用路径（无网络时必须优雅失败）')
    cred.enabled = True
    cred.base_url = 'https://api.qichacha.invalid'
    cred.timeout_sec = 2
    db.session.commit()
    check('is_configured=True', qcc_client.is_configured() is True)
    r = qcc_client.call('basic', '中国铁建股份有限公司', operator='tester', force_refresh=True)
    check('网络失败返回 fail 不抛异常', r.get('status') == 'fail', r.get('message')[:60])
    check('失败已记流水', ExtQueryLog.query.filter_by(ok=False).count() >= 1)
    check('流水消息已脱敏(无明文 key)',
          not any('abcdef1234567890KEY' in (l.message or '') for l in ExtQueryLog.query.all()))
    res2 = services.verify_subcontractor(prof, operator='tester', force_refresh=True)
    check('编排层网络失败仍返回结构化', res2['status'] == 'fail', res2.get('message', '')[:50])
    check('档案 qcc_status 已更新', prof.qcc_status == 'fail')

    print('\n[11] 配额限制')
    cred.daily_quota = 1
    db.session.commit()
    r = qcc_client.call('basic', '另一家公司', force_refresh=True)
    check('超配额被拦截', r.get('status') == 'quota', r.get('message'))
    cred.daily_quota = None
    db.session.commit()

    print('\n[12] 路由与端点完整性')
    eps = set(app.view_functions.keys())
    for ep in ['integration.datasource_list', 'integration.template_list',
               'integration.template_detail', 'integration.datasource_save',
               'integration.datasource_test', 'integration.field_batch_save',
               'integration.template_preview', 'subcontractor.profile_export_xlsx',
               'subcontractor.profile_export_batch']:
        check('端点存在 %s' % ep, ep in eps)

    print('\n[13] 页面渲染（模板语法）')
    with app.test_request_context('/'):
        from flask import render_template
        from flask_login import AnonymousUserMixin

        class FakeUser(AnonymousUserMixin):
            is_authenticated = True
            id = 1

            def has_permission(self, p):
                return True

            def get_data_scope(self):
                return 'all'

            def is_admin(self):
                return True

            def get_role_name(self):
                return '管理员'

            def get_visible_projects(self):
                return []

            def get_main_project(self):
                return None

            def __getattr__(self, name):
                # 模板/context processor 可能访问任意用户属性；未显式定义的方法/属性一律返回 None
                if name.startswith('__') and name.endswith('__'):
                    raise AttributeError(name)
                return None
        import flask_login.utils as flu
        orig = flu._get_user
        flu._get_user = lambda: FakeUser()
        try:
            html = render_template('integration/datasource_list.html',
                                   creds=ExtCredential.query.all(),
                                   endpoints=ExtEndpoint.query.all(),
                                   PROVIDERS={'qcc': '企查查开放平台'},
                                   today_calls=3, recent_logs=ExtQueryLog.query.limit(5).all(),
                                   cache_count=1)
            check('数据源页渲染', len(html) > 2000 and 'API Key' in html)
            html2 = render_template('integration/template_list.html',
                                    templates=ExportTemplate.query.all())
            check('模板列表页渲染', '导出模板管理' in html2)
            html3 = render_template('integration/template_detail.html',
                                    tpl=tpl, FIELD_SOURCES=FIELD_SOURCES)
            check('模板详情页渲染', '字段映射' in html3 and '整表保存' in html3)
            check('掩码 key 未泄露明文', 'abcdef1234567890KEY' not in html)
        finally:
            flu._get_user = orig

print('\n' + '=' * 60)
if FAIL:
    print('❌ 失败 %s 项：%s' % (len(FAIL), FAIL))
    sys.exit(1)
print('✅ 全部通过')
